import json
import logging
import asyncio
import os
import re
import time
import html
import difflib
from datetime import datetime, timedelta
from io import BytesIO
from aiohttp import web
from aiogram import Bot, Dispatcher, F, types, BaseMiddleware
from aiogram.filters import Command, CommandStart, StateFilter
from aiogram.types import (
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardMarkup,
    WebAppInfo,
    BotCommand,
    FSInputFile,
    BufferedInputFile
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage

# GOOGLE GENAI SDK ENTEGRASYONU
try:
    from google import genai
    from google.genai import types as genai_types
    HAS_GENAI = True
except ImportError:
    HAS_GENAI = False

# ==============================================================================
#                     CANLI YAPILANDIRMA & TEST MODU
# ==============================================================================
BOT_TOKEN = "9212674639:AAGVES4GiQRUlwrXfFy6iu4IsXKYo-4jzUY"
OMER_CHAT_ID = 6915998444
KAMIL_CHAT_ID = 35013468
WEBAPP_URL = "https://omeerkurt.github.io/eliar-depo/"
GEMINI_API_KEY = "AIzaSyAuUg0a1ew5V4aF1E-FTsdVinsMpXckOgQ"

# 🧪 TEST & SİMÜLASYON MODU (True: Kamil'e gitmez, gerçek stoklar ve Z-raporu değişmez)
TEST_MODE = True

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

ai_client = None
if HAS_GENAI and GEMINI_API_KEY:
    try:
        ai_client = genai.Client(api_key=GEMINI_API_KEY)
        logging.info("✅ Google GenAI (Gemini) istemcisi bağlandı.")
    except Exception as e:
        logging.warning(f"⚠️ Gemini bağlantı uyarısı: {e}")

class LoggingMiddleware(BaseMiddleware):
    async def __call__(self, handler, event, data):
        if isinstance(event, types.Message):
            u_name = event.from_user.first_name if event.from_user else "Bilinmeyen"
            logging.info(f"--- MESAJ: {u_name} (ID: {event.chat.id}) | {event.text or (event.document.file_name if event.document else '')}")
        return await handler(event, data)

dp.message.outer_middleware(LoggingMiddleware())

STOCKS_FILE = "stocks_db.json"
SALES_FILE = "weekly_sales.json"
TASKS_FILE = "active_tasks.json"
QUOTES_FILE = "submitted_quotes.json"
PDF_STORAGE_DIR = "stored_quotes_pdf"
os.makedirs(PDF_STORAGE_DIR, exist_ok=True)

# ----------------- 64 SAYFALIK KATALOĞUN TAM BİRLEŞİK LİSTESİ (170+ KALEM) -----------------
REAL_WAREHOUSE_STOCKS = [
    # Sayfa 2 - Güç ve Kontrol
    {"code": "91661", "name": "Sürücü VZAB2P2BAA V1000 3/2,2 kW", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "40642", "name": "PLC CPU: CJ1M-CPU21", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "40645", "name": "Güç kaynağı: CJ1W-PA202", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "40643", "name": "Dijital giriş modülü: CJ1W-ID211", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "40644", "name": "Dijital çıkış modülü: CJ1W-OC211", "cat": "ELEKTRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "90802", "name": "Güç Kaynağı: CQM1-PA203 (omron)", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 3 - PLC ve Pompalar
    {"code": "CQM1H-CPU21", "name": "Omron CQM1H-CPU21 Modül", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "CQM1-ID211", "name": "Omron CQM1-ID211 Input Modül", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41370", "name": "Plc Çıkış Modülü: CQM1-OC222 (Omron)", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41381", "name": "Rd96 22 Li Pompa Altı Paslanmaz Tava", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91514", "name": "Rd96 SLR1-40 AISI 316L Loblu Pompa", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 4 - Motor ve Bağlantılar
    {"code": "91410", "name": "Wat 3 faz Motor Q2E112M6C 2,2 kW 3 HP 1000 d/d", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91411", "name": "WAT Motor Ayağı (Q2E112M6C)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41436", "name": "Rd96 SLR1-40 Pompa Çıkış Bağlantı Rakor ve Dirsekleri", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "41435", "name": "Rd96 SLR1-40 Pompa Giriş Bağlantı Rakor ve Dirsekleri", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "96469", "name": "Rakor Contası - Gasket - Nw 40 Dın - EPDM", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "60002", "name": "Klemp Contası: EPDM (DN25) 1''- Ø33.7mm", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},

    # Sayfa 5 - SLR1-40 Pompa Parçaları
    {"code": "90338", "name": "SLR1-40 ASI 316L Salmastra Bileziği (Part No 13)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91128", "name": "SLR 1-40 Mec. Seal Si.Carbon/Si.Carbide (Part No 8-8A)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "91516", "name": "SLR1-40 NBR O-Ring (Part No 80D)", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "41538", "name": "SLR1-40 Cover Seal (Part No: 09) (AISI-304)", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "41407", "name": "SLR1-40 Pump Housing (Part No: 01) (CF-8M)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 6 - Lob ve Cıvatalar
    {"code": "41420", "name": "SLR1-40 Allen Screw (Part No: 51B) (A2)", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},
    {"code": "93041", "name": "SLR1-40 Brass Flanş Burcu (Part No 35)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "93035", "name": "SLR1-40 AISI 431 Wedge Lobe (Part No 2)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "60020", "name": "SLR1-40 AISI 316 Lobe Screw (Part No 25)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "91517", "name": "SLR1-40 NBR O-Ring (Part No 80)", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "41430", "name": "SLR1-40 O-ring (Part No: 80A) (EPDM)", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},

    # Sayfa 7 - Kapaklar ve Burçlar
    {"code": "41408", "name": "SLR1-40 Pump Cover (Part No: 03) (AISI-316)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41415", "name": "SLR1-40 Wing Nut (Part No: 45) (AISI 316)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41419", "name": "SLR1-40 Allen Screw (Part No: 51A) (8.8)", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},
    {"code": "41414", "name": "SLR1-40 Dragging Bushing (Part No: 17B)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},

    # Sayfa 8 - Dişliler ve Somunlar
    {"code": "41425", "name": "SLR1-40 Safety Nut (Part No: 62) (Steel)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41427", "name": "SLR1-40 Conical Tightening Ring (Part No: 65A) (Steel)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "41426", "name": "SLR1-40 Safety Washer (Part No: 63) (Steel)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "93039", "name": "SLR1-40 F 154 Motor Shaft Dişlisi (Part No 19)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "93040", "name": "SLR1-40 F-154 Tahrik Shaft Dişlisi (Part No 19A)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},

    # Sayfa 9 - Rulmanlar ve Burçlar
    {"code": "41413", "name": "SLR1-40 Drive Shaft Bushing (Part No:17) (ST-35)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41418", "name": "SLR1-40 Allen Screw (Part No: 51) (8.8)", "cat": "MEKATRONIK", "qty": 12, "loc": "ELIPS-DEPO"},
    {"code": "41411", "name": "SLR1-40 Bearing Stop Flange (Part No: 12A) (F-114)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41428", "name": "SLR1-40 Ball Bearing (Part No: 70) (Steel)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "41429", "name": "SLR1-40 Needle Bearing (Part No: 70A) (Steel)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},

    # Sayfa 10 - Şaftlar ve Keçeler
    {"code": "93043", "name": "SLR1-40 Steel Tahrik Shaft Mili Kaması (Part No 61A)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "93042", "name": "SLR1-40 Steel Motor Shaft Mili Kaması (Part No 61)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "93037", "name": "SLR1-40 AISI 329 Tahrik Shaft (Part No 05A)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "93036", "name": "SLR1-40 AISI 329 Motor Shaft (Part No 5)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "60044", "name": "SLR1-40 NBR Keçe (Part No 88)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41434", "name": "SLR1-40 Bleeder (Part No: 87) (Plastic)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},

    # Sayfa 11 - Gövde Elemanları
    {"code": "90946", "name": "SLR1-40 Filtre", "cat": "MEKATRONIK", "qty": 12, "loc": "ELIPS-DEPO"},
    {"code": "41432", "name": "SLR1-40 Oil Plug (Part No: 85) (Plastic)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "41409", "name": "SLR1-40 Support (Part No: 06) (GG-15)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41424", "name": "SLR1-40 Pin (Part No: 56) (F-522)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},

    # Sayfa 12 - Kelepçeler ve Kapaklar
    {"code": "41433", "name": "SLR1-40 Peephole (Part No: 86) (Plastic)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "41416", "name": "SLR1-40 Protector (Part No: 47A) (Plastic)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41417", "name": "SLR1-40 Screw (Part No: 50) (A2)", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "41533", "name": "SLR1-40 Part 18A (Gasket)", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "96541", "name": "23-35 Paslanmaz Hortum Kelepçesi", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "96540", "name": "13-19 mm Paslanmaz Hortum Kelepçesi", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},

    # Sayfa 13-17 - FIP25 Pompa Komple ve Elemanları
    {"code": "90951", "name": "Rd96 Fip25 Pompa + Motor+ Flanş (Takım)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41384", "name": "R96 Fip 25 Pompa Bağlantı Çıkış Rakoru ve Dirseği", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41383", "name": "R96 Fip 25 Pompa Bağlantı Giriş Rakoru ve Dirseği", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "90310", "name": "Fip25 1\"/1\" Esnek Çarklı Pompa (Motorsuz)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90952", "name": "Fip25 Pompa Gövde+Kapak+Oring", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90332", "name": "Fip25 Pompa Mili (01-46634) (Şaft Version 2)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90328", "name": "Fip25 Maşon (Sleeve 01-45214)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "90949", "name": "Fip25 Maşon Alyan Vida (0.0300.943)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "90947", "name": "Fip25 Salmastra Sacı (Washer 01-45782)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90321", "name": "Fip25 Saç Vidası (0.0150.001)", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "16073", "name": "Fip25 M8 Pul (Washer 0.0350.11)", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "90490", "name": "Fip25 Kapak Cıvatası Screw 4-20x8", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "91238", "name": "Fip25 O-ring (0-2173.523) FDA", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "90333", "name": "Fip25 Kapak Pimi (Screw Stud 01-46638)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "90948", "name": "Fip25 Bakalit Flanş (Motor Flange 01-24252)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91407", "name": "Fip25 Salmastra (09-46686-04) Si&Si", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "93008", "name": "Fip25 Pompa Gövdesi (SH-DIN)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90345", "name": "Fip25 Fan/Impeller (09-836S-7) FDA EPDM", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90325", "name": "Fip25 Pompa Kapağı (01-35850) Cover SH", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90320", "name": "Fip25 Kapak Somunu (Cap screw 01-46505)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "90950", "name": "Fip25 Motor Flanşı (Paslanmaz+Demir)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "60003", "name": "Rakor Contası - Gasket - Nw 25 Dın - EPDM", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "96112", "name": "3/4\" Siyah Madde Verici Kimyasal Hortum (MT)", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},

    # Sayfa 18-22 - Omal Vanalar, Aktüatörler, Switchbox
    {"code": "90827", "name": "Omal DG451H005 3 Yollu L Port 3/4\" (DN20) PN40", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50155", "name": "Omal Aktüatör: DA030401S DA30 30Nm Flanş F03-F05", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96467", "name": "Omal Aktüatör: DAN0030411S DA30 30Nm F03-F05", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91603", "name": "Omal Vana: V451H405 3 Yollu L Port 3/4\" (DN20)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91035", "name": "Omal (Tamir) Takımı: KGBV4505 3/4\" (DN20)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90547", "name": "3 Yollu \"L\" Port 3/4\" Vana Küresi (545-II-F05)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "20031", "name": "Tork LS102M Switch Box", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90475", "name": "Autonics Kapasitif Sensör: CR18-8DN", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50160", "name": "Pnömatik rakor: 57116 1/2\"-10mm dirsek pirinç", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "91285", "name": "P00142140 U2 Çek Valf", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50606", "name": "Nipel 2010 - 1/2\"", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "91472", "name": "Solenoid Valf: L03Eb13 1/2\" S00100400", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50247", "name": "Valf bobini: B00126625 FGC1 24VDC 6W", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50249", "name": "Valf soketi: M1TS2L01 B0122TH25 24V ledli", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "15401", "name": "Tee rakor: 4000 1/4\" D-D-D", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41439", "name": "1/4\" Paslanmaz Tek Taraflı Nipel", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91633", "name": "Omal VIP Vana: VNC10004 NBR 1/2\" (DN15) PN10", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91448", "name": "Pnömatik rakor: 57116 1/8\"-6mm dirsek pirinç", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "42595", "name": "1/2\"x Ø16mm Pirinç Altıköşe Hortum Ağzı", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "90861", "name": "Siyah Kauçuk Su Hortumu 16mm", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "90817", "name": "Omal Vana: DG424H405 3/4\" (DN20) PN65 15Nm", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96683", "name": "Omal Aktüatör: SRN0015401S SR15 15Nm F03-F05", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96097", "name": "DG424H405 Omal Aktüatör", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91598", "name": "Omal Küresel Vana: V424H405 3/4\" Paslanmaz", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91037", "name": "Omal (Tamir) Takımı: KGBV6005 3/4\" (DN20)", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "40082", "name": "Krohne Optiflux 5300 Magnetik Flowmeter", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 23-26 - Flowmetre, Diyafram ve Silindirler
    {"code": "91386", "name": "E/H DN25 (1/2\") mass Flowmetre (Promass40E15)", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91488", "name": "V424H405 Vana Küresi 3/4\" RD", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90567", "name": "Diyafram 66 Mm Teflon Takviyeli Diyafram", "cat": "MEKATRONIK", "qty": 9, "loc": "ELIPS-DEPO"},
    {"code": "90965", "name": "Manometre: G27-10-R1: AW10/AR10 için", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "40647", "name": "SMC CJPB15-10H6 Pin Silindir, Tek Etkili", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90715", "name": "Regülatör: AR10-M5H: 0-8 Bar", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50161", "name": "Nipel 2010 - M5", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50162", "name": "1021 Düz Erkek 6/4 M5", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "90828", "name": "Omal Vana: DG451H006 3Yollu L Port 1\" PN40", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90834", "name": "Omal Aktüatör: DA045402S Flanş: F03-F05", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96468", "name": "Omal Çift Etkili Aktüatör: DAN0045412S DA45", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91463", "name": "Eurotec Reducer - Kafa Adaptörü 11mm-9mm", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91604", "name": "V451H406 1\" Omal L Port Vana", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90548", "name": "3 Yollu \"L\" Port 1\" Vana Kuresi (545-II-F06)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91036", "name": "Kgbv4506 V451H006 Omal Vana Conta Takımı", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90832", "name": "Omal Aktüatör: DA0015411S DA15 15Nm F03", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "40634", "name": "BES0057: Balluff M12MI-NSC40B-BV03", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 27-31 - Çekvalfler, Nipeller, Rakorlar
    {"code": "96412", "name": "1\" > 3/4\" Polietilen Dst Çek Valf Adaptörü", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90878", "name": "Dst16-23 Cekvalf (Yaylı Çekvalf Paslanmaz 3/4\")", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90881", "name": "Dst Cek Valf Yayı", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "90880", "name": "DST Çekvalf Teflonu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90879", "name": "Dst Cek Valf Oring 007.59*2.62 EPDM", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "91284", "name": "Valf: P00112200 E1A/4: 3/2", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "14574", "name": "Nipel 2010 - 1/4\"", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "14541", "name": "Orantılı Nipel 2030 1/4\"-1/8\"", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "91332", "name": "Basınç Anahtarı: PE-1/8-1N: 1/8\" 0-8 bar", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "50219", "name": "Pnömatik rakor: 57116 1/4\"-10mm dirsek pirinç", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "54626", "name": "A03/0-24VDC B00123525: 5W 22mm bobin", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41441", "name": "57226 1/4\"x10 Döner Tee", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "91447", "name": "Pnömatik rakor: 57020 1/4\"-10mm düz erkek", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "50057", "name": "Pnömatik rakor: 57020 1/4\"-8mm: düz pirinç gövde", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "54612", "name": "Yıldız Başlı Cıvata: M4x30mm", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "50098", "name": "Manifold: SS5Y5-41-04-01F-Q: 4'lü", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50018", "name": "Valf: SY5140-5YO-Q: 5/2 tek bobin 24VDC", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "91594", "name": "Valf: SY5240-5YO-Q-5/2 Çift Bobin 1/8\" 24VDC", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "91590", "name": "Çek valf: U4: 1/4\"", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "15005", "name": "Kuyruklu dirsek: 5020 1/4\" erkek-dişi", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50549", "name": "Rediksiyon 2090 1/4\"<1/8\"", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "40085", "name": "Basınç anahtarı: PMN10A: 2-10 bar NO", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 32-37 - Valf Grupları, Raylar, Silindirler
    {"code": "91451", "name": "57116 RDD 1/4-8 DİRSEK AIGNEP", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "50023", "name": "KQ2S06-01AS: 1/8\"-6mm düz içten sıkmalı", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50273", "name": "SS5Y5-41-05-01F-Q 5'li Altlık", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50021", "name": "Valf soketi: DIN 43650C-C/Led 24V", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41442", "name": "RD96 MU_22 Krom Kaplı Ray: Ø25 / 180 mm", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41404", "name": "RD96 MU_33 Krom Kaplı Ray: Ø25 / 267 mm", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "96517", "name": "Rd96_MU Z Polietilen Z eksen Piston Dayama Takozu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50257", "name": "Tek Etlkili Pnömatik Silindir CD75E32-80-B (SMC)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96501", "name": "SMC C76-32PS Sızdırmazlık Kiti", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90766", "name": "RD96 Aliminyum Lama Montaj Parcası", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "16016", "name": "Paslanmaz Altı köşe Civata M5 x10 mm", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "16039", "name": "Paslanmaz M5 Yaylı Rondela", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "16034", "name": "Paslanmaz M5 Pul", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "41403", "name": "Havşa Başlı İmbus Civata M8 x 10", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},
    {"code": "90767", "name": "RD96 Cambaz Kolu Sacı", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91071", "name": "KQ2L04-01Ns Dirsek Rakor 1/8\" 4 Mm Nikel Kaplı", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "90762", "name": "Sensör tutucu: BM2-032: bant tipi 32mm çap", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90840", "name": "Piston sensörü: DC73L: reed 3m kablolu", "cat": "ELEKTRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41400", "name": "Rd96 MU Z ekseni Yataklama Pim leri 107*160 mm", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41401", "name": "Rd96 MU Z eksen Gijonları 292mm*M6", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41300", "name": "RD96-Mu Cambaz Pistonu Polietilen (Yeşil) Takoz", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41402", "name": "Rd96 MU Z ekseni Paslanmaz Yataklama Plakası", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91277", "name": "RD96-Mu Polietilen Ağızlık", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "54516", "name": "Rd96 Full Z Axis Equipment", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90768", "name": "Cambaz Kutusu Dolu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90769", "name": "RD96 Bos Cambaz Kutusu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90782", "name": "C85 Serisi Silindir C85C16 Eklem Yer Bağlantı", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "16032", "name": "Paslanmaz Altı köşe Civata M6 x10 mm", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "16041", "name": "Paslanmaz M6 Yaylı Rondela", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "16035", "name": "Paslanmaz M6 Pul", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "50260", "name": "Pnömatik Silindir CD85N16-50-B (SMC)", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "90759", "name": "Sensör tutucu: BJ2-016", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},

    # Sayfa 38-46 - Mafsallar, Kızaklar, Kilitlemeler
    {"code": "50138", "name": "Pnömatik rakor: KQ2L04-M5A: M5-4mm dirsek", "cat": "MEKATRONIK", "qty": 11, "loc": "ELIPS-DEPO"},
    {"code": "91086", "name": "Pnömatik T rakor: KQ2T04-M5A: orta M5-4mm", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "91045", "name": "SMC Küresel Mafsal: KJ6D M6*1", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "90770", "name": "RD96 Cambaz Mekanizması", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "16043", "name": "Paslanmaz İmbus M5 Civata M5x20 mm", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "16030", "name": "Paslanmaz Altı köşe Civata M5 x20 mm", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "90765", "name": "RD96 Cambaz Kolu", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "11906", "name": "Altı Köşe Somun M6", "cat": "MEKATRONIK", "qty": 20, "loc": "ELIPS-DEPO"},
    {"code": "12302", "name": "Fiberli Kör Somun M6", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},
    {"code": "41444", "name": "Krom Ray: Ø25 / 1350 mm RD96-MU22-MU33 x ekseni", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91303", "name": "Paslanmaz Araba 4 Delikli Krom Kaplı Bilyalı Rexroth", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50263", "name": "Silindir: CD85N25-400C-B Çelik (SMC)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91044", "name": "SMC Küresel Mafsal: KJ10D M10*1,25", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "54510", "name": "Pnömatik rakor: KQ2VS04-01AS: 1/8\"-4mm dirsek", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90761", "name": "Sensör tutucu: BM2-025: bant tipi", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "50262", "name": "Silindir: CD85N25-300C-B-XC6A (SMC)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41399", "name": "Pnömatik rakor: KQ2U04-00: 4mm Y bağlantı", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50259", "name": "Silindir CD85N16-200-B (SMC)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "50258", "name": "Silindir CD85N16-100-B (SMC)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "50261", "name": "Silindir 80 mm Strok CD85N16-80-B (SMC)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91269", "name": "RD96 X Ekseni mekanik yatak Kilidi", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91270", "name": "RD96 Y Ekseni mekanik yatak Kilidi", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50278", "name": "Manifold: SS5YJ3-S41-02-M5F-Q (3243 için 2'li Altlık)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50293", "name": "Valf: SYJ3243-5LOU-Q: 5/2 çift bobin", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "90725", "name": "Hız ayar valfi: ASN2-02: 1/8\" susturuculu", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},

    # Sayfa 47-52 - Valfler ve Donanımlar
    {"code": "50040", "name": "Dirsek Rakor 1/8x6mm / KQ2L06-01AS", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "41396", "name": "12'li Hareketli Kablo X3-X4 Pistonu Valf Kablosu", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41398", "name": "Soket ucu SY100-30-A 00044321", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "54630", "name": "SYJ3243 2'li full valf grubu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50270", "name": "Manifold: SS3Y1-S41-04-M5-Q: 4'lü Altlık", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50287", "name": "Valf: SY114-5LOU-Q", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "91066", "name": "Pnömatik rakor: KQ2H06-M5A: M5-6mm altıgen", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "91064", "name": "Pnömatik rakor: KQ2H04-M5A: M5-4mm düz", "cat": "MEKATRONIK", "qty": 13, "loc": "ELIPS-DEPO"},
    {"code": "41457", "name": "Kör Tapa: M-5P: M5x0,8 (SMC)", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "50238", "name": "Susturucu: AN120-M5 (SMC)", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50156", "name": "Pnömatik T rakor: KQ2T06-00A: 6mm", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91043", "name": "Valf Kablosu: SY114 için (EM)", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50269", "name": "Manifold: SS3Y1-S41-02-M5-Q: 2'li Altlık", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50289", "name": "Valf: SY3220-5LOU-M5-Q: çift bobin", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "50239", "name": "3220 İçin Braket", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50291", "name": "Valf: SYJ3143-5LOU-Q", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "90781", "name": "SMC C85-25Ps Piston Keçesi X3-X4", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},

    # Sayfa 53-57 - Tablalar, Vanalar, Hortumlar ve Soketler
    {"code": "41407P", "name": "RD96_MU22 PVC Tabla", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96404", "name": "RD96-MU22 Paslanmaz Çelik Tabla", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41406", "name": "RD96_MU33 PVC Tabla", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41405", "name": "RD96-MU33 Paslanmaz Çelik Tabla (304)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90484", "name": "3/4\" Kimyasal Vana / 3/4'' Kür.Vana+Elcik+Adaptör", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91490", "name": "3/4 RD Vana Adaptör Bağlantı O-ringi", "cat": "MEKATRONIK", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "91489", "name": "3/4' RD Vana Conta Takımı", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90930", "name": "EPDM Lastik Conta 25*40 Kare Oring", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "91495", "name": "3/4\" Seffaf Kimyasal Hortum", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "54661", "name": "Rd96 Mu 22 Y eksen Kolu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "54662", "name": "Rd96 Mu 33 Y eksen Kolu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41437", "name": "Flowmetre Dönüşüm Flanşı (Micromotion>E/Hauser)", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90568", "name": "66 Mm Teflon Takviyeli Diyafram Pulu", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "50014", "name": "Susturucu: AN15-02: 1/4\"", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "91782", "name": "Rd96-33 Hareketli Kablo (EM)", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91271", "name": "RD96-22 Hareketli Kablo (EM)", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91268", "name": "Rd96 Mu Fırca", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90793", "name": "Loblu Pompa Conta (Teflon Giydirme)", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "90324", "name": "SLR1-40 Lob Anahtarı", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "16044", "name": "Yıldız Tornavida başlı M4 x 5 mm vida", "cat": "MEKATRONIK", "qty": 15, "loc": "ELIPS-DEPO"},
    {"code": "14847", "name": "Klemp Kelepçesi: 1\" - Ø33.7mm Paslanmaz (316)", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41290", "name": "Soket muhafazası: HAN24 B-AGG-QB", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41287", "name": "Socket Cover HDC HB 24TSVU 1/21 P", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41288", "name": "Soket: HDC-HE-24BS: 24 pin dişi", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "41289", "name": "Soket: HDC-HE-24SS: 24 pin erkek", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},

    # Sayfa 58-63 - Purge Valfleri, Hortumlar ve Regülatörler
    {"code": "54629", "name": "SY114 4'lü Full Valf Grubu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "54628", "name": "SY114 2'li Full Valf Grubu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "54627", "name": "Dst komple purge valfi komple", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50024", "name": "Hava hortumu: TU0604BU: 6mm dış çap; mavi", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "50025", "name": "Hava hortumu: TU0806BU: 8mm dış çap; mavi", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "91586", "name": "Hava hortumu: TU1065BU: 10mm dış çap; mavi", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "50027", "name": "Hava hortumu: TU0425BU: 4mm dış çap mavi", "cat": "MEKATRONIK", "qty": 25, "loc": "ELIPS-DEPO"},
    {"code": "92038", "name": "Regülatör: AW30-F03DE-B: 3/8\" filtreli manometreli", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "54620", "name": "G36-P10-01-X30 Manometre", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "14577", "name": "Nipel: 2010 3/8\"", "cat": "MEKATRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "14317", "name": "Pnömatik Tee rakor: 4000 3/8\" D-D-D", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "96449", "name": "Orantılı nipel: 2030 1/4\"-3/8\"", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "91950", "name": "Çek valf: U4: 1/4\"", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "96219", "name": "Küresel vana: 6310 3/8\" erkek dişi", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "91453", "name": "Pnömatik rakor: 57116 3/8\"-10mm: dirsek pirinç", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "50008", "name": "Tespit Elemanı AR30 / AW30 / AR30P-270AS", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50010", "name": "Tespit Elemanı : AW23P-270AS", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50046", "name": "Regülatör: AW20-F02C-B (SMC)", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50054", "name": "Küresel vana: 6310 1/4\" erkek dişi", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},

    # Excel Tablosundan Gelen Kontrolörler ve Ekstra Parçalar
    {"code": "Pd320", "name": "Eliar Pd320 Boya Kontrolörü", "cat": "KONTROL", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "PD330", "name": "Eliar PD330 Boya Kontrolörü", "cat": "KONTROL", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "Aim8-A", "name": "Aim8-A Analog Giriş Modülü", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "Aim8-B", "name": "Aim8-B Analog Giriş Modülü", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "Dim32", "name": "Dim32 Dijital Giriş Modülü", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "Dom16-PLC", "name": "Dom16-PLC Dijital Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 6, "loc": "ELIPS-DEPO"},
    {"code": "P44", "name": "P44 Panel Kartı", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "Pd300", "name": "Eliar Pd300 Boya Kontrol Cihazı", "cat": "KONTROL", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "Pd340", "name": "Eliar Pd340 Boya Kontrolörü", "cat": "KONTROL", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "Pd350", "name": "Eliar Pd350 Boya Kontrolörü", "cat": "KONTROL", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "S400", "name": "Eliar S400 Kontrol Cihazı", "cat": "KONTROL", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "S42", "name": "Eliar S42 Kontrol Cihazı", "cat": "KONTROL", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "Umf-ILETISIM20", "name": "Umf-ILETISIM20 İletişim Kablosu", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "TBBPLC-F", "name": "TBBPLC-F Modülü", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "Edc T7000", "name": "Eliar Edc T7000 Kontrolör", "cat": "KONTROL", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "Edc M7700", "name": "Eliar Edc M7700 Kontrolör", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "Edc T710", "name": "Eliar Edc T710 Kontrolör", "cat": "KONTROL", "qty": 9, "loc": "ELIPS-DEPO"},
    {"code": "Edc T712", "name": "Eliar Edc T712 Kontrolör", "cat": "KONTROL", "qty": 8, "loc": "ELIPS-DEPO"},
    {"code": "Powerfail", "name": "Powerfail Modülü", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "P42", "name": "P42 Panel Kartı", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "P300", "name": "P300 Panel Kartı", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "P40", "name": "P40 Panel Kartı", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "Edc T7701EX", "name": "Eliar Edc T7701EX Kontrolör", "cat": "KONTROL", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "Edc T7700", "name": "Eliar Edc T7700 Kontrolör", "cat": "KONTROL", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "T7700", "name": "Eliar T7700 Cihazı", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "P43", "name": "P43 Panel Kartı", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "AOM8", "name": "AOM8 Analog Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 7, "loc": "ELIPS-DEPO"},
    {"code": "T7701 EX G", "name": "Eliar T7701 EX G Kontrolör", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "T7RACK 6 PLC", "name": "T7RACK 6 PLC Taşıyıcı", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "T7RACK 10 PLC", "name": "T7RACK 10 PLC Taşıyıcı", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-POWER 30", "name": "T7PAC-POWER 30 Güç Kaynağı", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-POWER 20", "name": "T7PAC-POWER 20 Güç Kaynağı", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-DIM32V5", "name": "T7PAC-DIM32V5 Giriş Modülü", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-DOM32V5", "name": "T7PAC-DOM32V5 Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-DIM16V5", "name": "T7PAC-DIM16V5 Giriş Modülü", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-AIO88SV1", "name": "T7PAC-AIO88SV1 Giriş/Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-AI8SV1", "name": "T7PAC-AI8SV1 Analog Giriş Modülü", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-AO8SV1", "name": "T7PAC-AO8SV1 Analog Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "PANELSET-T7701 ex", "name": "PANELSET-T7701 ex Ön Panel Seti", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "PANELSET-T7700", "name": "PANELSET-T7700 Ön Panel Seti", "cat": "KONTROL", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-T7CPURIO", "name": "T7PAC-T7CPURIO İşlemci Modülü", "cat": "KONTROL", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "T7PAC-DOM16V5", "name": "T7PAC-DOM16V5 Çıkış Modülü", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "41835", "name": "Apexco Salmastra 7003TTG-20 TUC/TUC.G/Vit/G45", "cat": "MEKATRONIK", "qty": 5, "loc": "ELIPS-DEPO"},
    {"code": "1013", "name": "Vana TORK 3/4\" 3 yollu + aktüatör", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "1015", "name": "SLD Kanal Vanası Full 3/4\" TORK 2 yollu", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "5022", "name": "Salmastra (Geylani ile geldi)", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "30010", "name": "EBARA 15MM SIL/SIL/VITON SALMASTRA", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "30076", "name": "SANTOPRENE DIAPHRAGM FOR POMKAR PK10", "cat": "MEKATRONIK", "qty": 4, "loc": "ELIPS-DEPO"},
    {"code": "40597", "name": "BES005M BALLUFF Sensör", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "41316", "name": "Lab120 encoder", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "42184", "name": "PIPETLEME SIRINGA UCU", "cat": "MEKATRONIK", "qty": 200, "loc": "ELIPS-DEPO"},
    {"code": "50090", "name": "Tork Çift Etkili Aktüatör", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "50227", "name": "KISICILI RAKOR: AS2201F-01-04SA 1/8\"-4MM", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50242", "name": "KISICILI RAKOR: AS2002F-06 6MM HAT TIPI", "cat": "MEKATRONIK", "qty": 7, "loc": "ELIPS-DEPO"},
    {"code": "50292", "name": "VALF: SYJ3240-5LOU-Q", "cat": "ELEKTRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "50294", "name": "VALF: SYJ7120-5DZ-01F-Q", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "50295", "name": "VALF: SYJ7220-5DZ-01F-Q", "cat": "ELEKTRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "52156", "name": "PNOMATIK T RAKOR: KQ2T06-00A: 6MM", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "54654", "name": "LOP POMPA FILTRE ORINGI 30*3 MM EPDM", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "90416", "name": "SOLENOID VALF: 00125317 6013-A02 24VDC", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "90501", "name": "SMC C76-32PS PISTON KECESI", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90726", "name": "HIZ AYAR VALFI: ASN2-02 1/4\" Susturuculu", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90771", "name": "RD96 CAMBAZ MEKANIZMASI", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "91074", "name": "PNOMATIK RAKOR: KQ2L6M5N Dirsek", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "91088", "name": "T RAKOR: KQ2T08-06 8-6MM", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91406", "name": "FIP25 SALMASTRA", "cat": "MEKATRONIK", "qty": 3, "loc": "ELIPS-DEPO"},
    {"code": "91502", "name": "LAB120 mikser kayışı", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "96496", "name": "Omal Çift Etkili Aktüatör: DAN0015411S DA15", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "96515", "name": "DST Polietilen Z Eksen dayama takozu", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "96516", "name": "DST Z Ekseni Polietilen Ağızlık", "cat": "MEKATRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "96570", "name": "Ø19mm Şeffaf Telli Hortum (Mavi)", "cat": "MEKATRONIK", "qty": 10, "loc": "ELIPS-DEPO"},
    {"code": "C-137161", "name": "Tuz rafinasyon filtre kapağı", "cat": "MEKATRONIK", "qty": 2, "loc": "ELIPS-DEPO"},
    {"code": "90620", "name": "PH Sensörü: SPH-1-DJ-6M (Seko)", "cat": "ELEKTRONIK", "qty": 0, "loc": "ELIPS-DEPO"},
    {"code": "SPU-PANO", "name": "Saha panosu (SPU su yok)", "cat": "ELEKTRONIK", "qty": 1, "loc": "ELIPS-DEPO"},
    {"code": "RA032DA", "name": "RA032DA09036000 Aktüatör", "cat": "MEKATRONIK", "qty": 1, "loc": "ELIPS-DEPO"}
]

class ClientHistoryState(StatesGroup):
    waiting_client_name = State()

class CustomReportState(StatesGroup):
    waiting_date_range = State()

class QuoteBuilderState(StatesGroup):
    waiting_firm = State()
    waiting_item_query = State()
    waiting_qty = State()
    confirm_more = State()

class CompletedQuoteSearchState(StatesGroup):
    waiting_firm_query = State()

def normalize_tr(text: str) -> str:
    if not text: return ""
    m = {'ı':'i','İ':'i','I':'i','ş':'s','Ş':'s','ğ':'g','Ğ':'g','ü':'u','Ü':'u','ö':'o','Ö':'o','ç':'c','Ç':'c'}
    res = str(text).lower()
    for k, v in m.items(): res = res.replace(k, v)
    return res

def load_data(file_path, default):
    if not os.path.exists(file_path):
        save_data(file_path, default)
        return default
    try:
        with open(file_path, "r", encoding="utf-8") as f: return json.load(f)
    except Exception: return default

def save_data(file_path, data):
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def sync_all_stocks_to_db():
    """Tüm 170+ katalog ve excel parçasını diske zorunlu eşitler"""
    existing = load_data(STOCKS_FILE, [])
    existing_map = {s["code"].lower(): s for s in existing}
    
    # Eksik olanları ekle
    for p in REAL_WAREHOUSE_STOCKS:
        code_l = p["code"].lower()
        if code_l not in existing_map:
            existing.append(p)
            existing_map[code_l] = p
            
    save_data(STOCKS_FILE, existing)
    logging.info(f"✅ TÜM KATALOG EŞİTLENDİ: {len(existing)} Kalem Parça Veritabanında Aktif!")

load_stocks = lambda: load_data(STOCKS_FILE, REAL_WAREHOUSE_STOCKS)
load_sales = lambda: load_data(SALES_FILE, [])
load_quotes = lambda: load_data(QUOTES_FILE, [])
save_quotes = lambda d: save_data(QUOTES_FILE, d)

# ----------------- TEST MODU GÜVENLİK KİLİTLERİ -----------------
def save_stocks(d):
    if TEST_MODE:
        logging.info("🧪 TEST MODU: save_stocks engellendi (Stoklar sabit kaldı).")
        return
    save_data(STOCKS_FILE, d)

def save_sales(d):
    if TEST_MODE:
        logging.info("🧪 TEST MODU: save_sales engellendi (Z-Raporu temiz kaldı).")
        return
    save_data(SALES_FILE, d)

def get_catalog_pdf_file():
    for name in ["Katalog.pdf", "katalog.pdf", "CATALOG.PDF", "catalog.pdf", "rd96.pdf", "RD96.pdf"]:
        if os.path.exists(name): return name
    for root, dirs, files in os.walk("."):
        for fn in files:
            lower = fn.lower()
            if lower.endswith(".pdf") and ("katalog" in lower or "catalog" in lower or "rd96" in lower):
                return os.path.join(root, fn)
    return None

# ==============================================================================
#      KUSURSUZ SAYFALAR ARASI AKILLI KATALOG KIRPICI (CROSS-PAGE STITCHER)
# ==============================================================================
def stitch_pixmaps_vertically(pix_header, pix_body) -> bytes:
    try:
        from PIL import Image
        import io
        img1 = Image.open(io.BytesIO(pix_header.tobytes("png")))
        img2 = Image.open(io.BytesIO(pix_body.tobytes("png")))

        target_w = max(img1.width, img2.width)
        if img1.width != target_w:
            img1 = img1.resize((target_w, int(img1.height * (target_w / img1.width))))
        if img2.width != target_w:
            img2 = img2.resize((target_w, int(img2.height * (target_w / img2.width))))

        combined = Image.new("RGB", (target_w, img1.height + img2.height), (255, 255, 255))
        combined.paste(img1, (0, 0))
        combined.paste(img2, (0, img1.height))

        out = io.BytesIO()
        combined.save(out, format="PNG")
        return out.getvalue()
    except Exception as e:
        logging.warning(f"Resim birleştirme hatası: {e}")
        return pix_body.tobytes("png")

def crop_slot_from_page(page, slot: int) -> bytes:
    pw, ph = page.rect.width, page.rect.height
    col = 0 if slot in [1, 3, 5] else 1
    row = 0 if slot in [1, 2] else (1 if slot in [3, 4] else 2)

    x0 = 8 if col == 0 else (pw / 2.0 + 2)
    x1 = (pw / 2.0 - 2) if col == 0 else (pw - 8)

    row_bounds = [
        (25, 285),
        (288, 550),
        (550, 815)   # Slot 5 & 6 (90567 Diyafram)
    ]
    y0, y1 = row_bounds[row]

    try:
        import fitz
        clip_rect = fitz.Rect(max(0, x0), max(0, y0), min(pw, x1), min(ph, y1))
        pix = page.get_pixmap(clip=clip_rect, dpi=200)
        return pix.tobytes("png")
    except Exception as e:
        logging.error(f"Kırpma hatası: {e}")
        return None

def extract_part_image_from_pdf(code: str, name: str = "") -> bytes:
    pdf_path = get_catalog_pdf_file()
    if not pdf_path: return None
    try:
        import fitz
        doc = fitz.open(pdf_path)
        code_str = str(code).strip()

        # 1. Önce PDF içindeki gerçek vektör metin aramasını yap
        for p_idx, page in enumerate(doc):
            rects = (
                page.search_for(f"({code_str})")
                or page.search_for(f" {code_str} ")
                or page.search_for(code_str)
            )
            if rects:
                pw, ph = page.rect.width, page.rect.height
                r = rects[0]
                col = 0 if ((r.x0 + r.x1) / 2.0) < (pw / 2.0) else 1
                x0 = 8 if col == 0 else (pw / 2.0 + 2)
                x1 = (pw / 2.0 - 2) if col == 0 else (pw - 8)

                # DURUM 1: Başlık önceki sayfanın dibinde, resim bu sayfada (Örn: Pin Silindir 40647)
                if r.y0 < 130 and p_idx > 0:
                    p_prev = doc[p_idx - 1]
                    rect_header = fitz.Rect(x0, ph - 62, x1, ph - 20)
                    pix_header = p_prev.get_pixmap(clip=rect_header, dpi=200)
                    rect_body = fitz.Rect(x0, max(15, r.y0 - 15), x1, min(ph - 15, r.y0 + 245))
                    pix_body = page.get_pixmap(clip=rect_body, dpi=200)
                    return stitch_pixmaps_vertically(pix_header, pix_body)

                # DURUM 2: Başlık bu sayfanın dibinde, resim sonraki sayfada
                if r.y0 > ph - 100 and p_idx < len(doc) - 1:
                    p_next = doc[p_idx + 1]
                    rect_header = fitz.Rect(x0, max(15, r.y0 - 15), x1, ph - 20)
                    pix_header = page.get_pixmap(clip=rect_header, dpi=200)
                    rect_body = fitz.Rect(x0, 25, x1, 285)
                    pix_body = p_next.get_pixmap(clip=rect_body, dpi=200)
                    return stitch_pixmaps_vertically(pix_header, pix_body)

                # DURUM 3: Standart Tek Sayfa (Örn: Diyafram 90567)
                y0 = max(10, r.y0 - 32)
                y1 = min(ph - 10, y0 + 268)
                clip_rect = fitz.Rect(x0, y0, x1, y1)
                pix = page.get_pixmap(clip=clip_rect, dpi=200)
                return pix.tobytes("png")

        # 2. Dinamik bulunamazsa güncellenmiş koordinat haritasından çek
        if code_str in CATALOG_GRID_MAP:
            page_no, slot = CATALOG_GRID_MAP[code_str]
            p_idx = page_no - 1
            if p_idx < len(doc):
                return crop_slot_from_page(doc[p_idx], slot)

        return None
    except Exception as e:
        logging.error(f"Görsel çıkarma hatası: {e}")
        return None

# ==============================================================================
#      KUSURSUZ ELIAR TEKLİF PDF OKUYUCU (YEREL TABLO + GEMINI DUAL ENGINE)
# ==============================================================================
def parse_eliar_pdf_locally(pdf_path: str, original_filename: str = "") -> dict:
    full_text = ""
    items = []
    firm = None
    ref_no = "BELİRTİLMEDİ"
    date_str = datetime.now().strftime("%d.%m.%Y")
    grand_total = "-"
    currency = "EUR"
    is_nilay = False

    try:
        import fitz
        doc = fitz.open(pdf_path)
        for page in doc: full_text += page.get_text() + "\n"
    except Exception: pass

    # 1. Müşteri Adı Tespiti
    m_firm = re.search(r"Müşteri\s*Adı\s*:\s*([^\n\r]+)", full_text, re.IGNORECASE)
    if m_firm:
        val = m_firm.group(1).strip()
        val = re.split(r"(?:İlgili\s*Kişi|İlgili|Tel|Adres)", val, flags=re.IGNORECASE)[0].strip()
        if len(val) >= 3 and val.upper() != "ADI:":
            firm = val.upper()

    if not firm or firm == "ADI:":
        clean_fn = re.sub(r"(fiyat|teklif|teklifi|\.pdf|\.jpg|\.png)", "", original_filename, flags=re.IGNORECASE).strip()
        clean_fn = re.sub(r"^[_\-\s]+|[_\-\s]+$", "", clean_fn)
        if len(clean_fn) >= 3:
            firm = clean_fn.upper()
        else:
            firm = "BİLİNMEYEN FİRMA"

    # 2. Ref No & Tarih Tespiti
    m_ref = re.search(r"Ref\s*No\s*:\s*([A-Za-z0-9\-]+)", full_text, re.IGNORECASE)
    if m_ref: ref_no = m_ref.group(1).strip()

    m_date = re.search(r"Teklif\s*Tarihi\s*:\s*([\d\.\/\-]+)", full_text, re.IGNORECASE)
    if m_date: date_str = m_date.group(1).strip()

    if "nilay" in full_text.lower():
        is_nilay = True

    # 3. Toplam Tutar & Para Birimi
    m_tot = re.search(r"TOPLAM\s*:\s*([\d\.,]+)\s*([€$£]|EUR|USD|TL)?", full_text, re.IGNORECASE)
    if m_tot:
        grand_total = m_tot.group(1).strip()
        curr_symbol = m_tot.group(2) or "€"
        currency = "EUR" if curr_symbol == "€" else ("USD" if curr_symbol == "$" else curr_symbol)

    # 4. Tablo ve Parçaları Çekme
    try:
        import pdfplumber
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables() or []
                for table in tables:
                    for row in table:
                        if not row or len(row) < 3: continue
                        clean_row = [str(c).replace('\n', ' ').strip() if c else '' for c in row]
                        row_str = " ".join(clean_row)

                        if "Ürün Kodu" in row_str or "Birim Fiyat" in row_str or "TOPLAM" in row_str:
                            continue

                        c0 = clean_row[0]
                        c1 = clean_row[1] if len(clean_row) > 1 else ""

                        if re.match(r"^\d{4,7}$", c0) or (len(c0) >= 3 and re.match(r"^[A-Za-z0-9\-_./]+$", c0)):
                            q = 1
                            for cell in clean_row[2:]:
                                digits = re.sub(r"[^\d]", "", cell)
                                if digits and 0 < int(digits) < 10000:
                                    q = int(digits)
                                    break
                            up = clean_row[3] if len(clean_row) > 3 else "-"
                            tp = clean_row[4] if len(clean_row) > 4 else "-"
                            items.append({
                                "code": c0,
                                "name": c1,
                                "qty": q,
                                "unit_price": up,
                                "total_price": tp
                            })
    except Exception as e:
        logging.warning(f"pdfplumber ayrıştırma: {e}")

    if not items:
        for line in full_text.splitlines():
            m_item = re.search(r"(\b\d{4,7}\b)\s+([A-Za-z0-9\s\/\-\.\(\)\"\,\*]+?)\s+(\d+)\s+([\d\.,]+\s*€?)", line)
            if m_item:
                items.append({
                    "code": m_item.group(1).strip(),
                    "name": m_item.group(2).strip(),
                    "qty": int(m_item.group(3).strip()),
                    "unit_price": m_item.group(4).strip(),
                    "total_price": "-"
                })

    return {
        "firm": firm,
        "ref_no": ref_no,
        "date": date_str,
        "currency": currency,
        "grand_total": grand_total,
        "items": items,
        "is_nilay": is_nilay
    }

async def parse_quote_pdf_smart(pdf_path: str, original_filename: str = "") -> dict:
    local_data = parse_eliar_pdf_locally(pdf_path, original_filename)
    if local_data["items"] and local_data["firm"] != "BİLİNMEYEN FİRMA":
        return local_data

    if ai_client:
        try:
            loop = asyncio.get_event_loop()
            prompt = (
                "Teklif belgesinden şu alanları saf JSON dön:\n"
                '{"firm":"...","ref_no":"...","date":"...","currency":"EUR","grand_total":"...","items":[{"code":"...","name":"...","qty":1,"unit_price":"...","total_price":"..."}]}'
            )
            with open(pdf_path, "rb") as f: file_bytes = f.read()
            for m_name in ["gemini-2.0-flash", "gemini-1.5-flash"]:
                try:
                    response = await loop.run_in_executor(
                        None,
                        lambda: ai_client.models.generate_content(
                            model=m_name,
                            contents=[prompt, genai_types.Part.from_bytes(data=file_bytes, mime_type="application/pdf")]
                        )
                    )
                    raw = response.text.replace("```json", "").replace("```", "").strip()
                    parsed = json.loads(raw)
                    if parsed.get("items"):
                        parsed["is_nilay"] = local_data.get("is_nilay", False)
                        return parsed
                except Exception:
                    pass
        except Exception as e:
            logging.error(f"Gemini API parse hatası: {e}")

    return local_data

def calculate_similarity(query: str, part: dict) -> float:
    q_norm = normalize_tr(query.strip())
    name_norm = normalize_tr(part["name"])
    code_norm = normalize_tr(part["code"])
    if q_norm == code_norm: return 1.0
    if q_norm in code_norm: return 0.95
    if q_norm in name_norm: return 0.88
    q_words, name_words = q_norm.split(), name_norm.split()
    word_scores = [max([difflib.SequenceMatcher(None, qw, nw).ratio() for nw in name_words] or [0]) for qw in q_words]
    avg_word = sum(word_scores) / len(word_scores) if word_scores else 0.0
    return max(avg_word, difflib.SequenceMatcher(None, q_norm, name_norm).ratio())

def find_best_matches(query: str):
    stocks = load_stocks()
    scored = [(p, calculate_similarity(query, p)) for p in stocks if calculate_similarity(query, p) > 0.30]
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored

# ----------------- TEST MODUNA DUYARLI BİLDİRİM -----------------
async def notify_both(text: str, reply_markup=None):
    recipients = {OMER_CHAT_ID} if TEST_MODE else {OMER_CHAT_ID, KAMIL_CHAT_ID}
    if TEST_MODE:
        text += "\n\n🧪 <i>[TEST MODU: Kamil'e bildirim gitmedi, gerçek stoklar ve Z-raporu korunuyor.]</i>"

    for cid in recipients:
        try: await bot.send_message(chat_id=cid, text=text, parse_mode="HTML", reply_markup=reply_markup)
        except Exception as e: logging.warning(f"İletilemedi ({cid}): {e}")

# ----------------- GÜVENLİ Z-RAPORU MOTORU -----------------
def parse_date_safe(date_str: str) -> datetime:
    if not date_str: return None
    date_str = str(date_str).strip()
    try:
        if "T" in date_str: return datetime.fromisoformat(date_str.split(".")[0])
    except Exception: pass
    for fmt in ["%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"]:
        try: return datetime.strptime(date_str, fmt)
        except Exception: pass
    return None

def filter_sales_by_range(start_dt: datetime, end_dt: datetime):
    sales = load_sales()
    res = []
    for s in sales:
        if s.get("status") != "APPROVED": continue
        d = parse_date_safe(s.get("timestamp") or s.get("date"))
        if d and (start_dt <= d <= end_dt):
            res.append(s)
    return res

def build_report_text(sales_data, report_name="SATIŞ Z-RAPORU"):
    now = datetime.now()
    if not sales_data:
        return (
            f"📊 <b>ELİPS DEPO {report_name}</b>\n"
            f"📅 Tarih: {now.strftime('%d.%m.%Y %H:%M')}\n"
            "═══════════════════════════════\n"
            "ℹ️ <i>Bu dönem için onaylanmış parça satışı/çıkışı bulunmamaktadır.</i>"
        )

    item_totals = {}
    plants = set()
    total_qty = 0

    for sale in sales_data:
        plants.add(sale.get("firm", "Bilinmeyen"))
        for itm in sale.get("items", []):
            code = itm.get("code", "KODSUZ")
            name = itm.get("name", "Ürün")
            q = int(itm.get("qty", 1))
            total_qty += q
            if code not in item_totals:
                item_totals[code] = {"name": name, "qty": 0}
            item_totals[code]["qty"] += q

    text = (
        f"📊 <b>ELİPS DEPO {report_name}</b>\n"
        f"📅 Rapor Tarihi: {now.strftime('%d.%m.%Y %H:%M')}\n"
        "═══════════════════════════════\n"
        f"💰 <b>GENEL SEVKİYAT ÖZETİ:</b>\n"
        f"• Toplam Satılan Parça: <b>{total_qty} Adet</b>\n"
        f"• Hizmet Verilen Fabrika: <b>{len(plants)} Tesis</b>\n"
        f"• Tamamlanan İşlem: <b>{len(sales_data)} Hareket</b>\n"
        "═══════════════════════════════\n\n"
        "<b>📦 SATILAN PARÇALAR (KONSOLİDE LİSTE):</b>\n"
    )

    for idx, (code, data) in enumerate(item_totals.items(), 1):
        text += f"• <code>[{code}]</code> {html.escape(data['name'][:30])} ➔ <b>{data['qty']} Adet</b>\n"

    text += "\n═══════════════════════════════\n<b>🏭 TESİS ÇIKIŞLARI:</b>\n"
    for s in sales_data[:10]:
        depo_tag = "🇹🇷 [TR]" if s.get("source") == "TR_DEPO" else "🇺🇿 [ELİPS]"
        item_summary = ", ".join([f"{i.get('code')} ({i.get('qty')} adt)" for i in s.get("items", [])])
        text += f"• <code>{s.get('date', '')[:10]}</code> {depo_tag} <b>{html.escape(s.get('firm', ''))}</b>\n  └ {html.escape(item_summary)}\n"

    if len(sales_data) > 10:
        text += f"\n<i>...ve {len(sales_data) - 10} hareket daha (Aşağıdan Excel veya PDF ile tam dökümü alabilirsiniz).</i>\n"

    text += "\n📂 <i>İndirmek istediğiniz dosya formatını seçin:</i>"
    return text

def generate_excel_z_report(sales_data, title="Z_Raporu"):
    fn = f"Elips_{title}_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Satislar"
        headers = ["Sıra", "Tarih", "Kaynak Depo", "Ref No", "Müşteri / Fabrika", "Ürün Kodu", "Ürün Adı", "Adet"]
        ws.append(headers)
        h_fill = PatternFill(start_color="161F30", end_color="161F30", fill_type="solid")
        h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border = Border(*(Side(style='thin', color='DDDDDD') for _ in range(4)))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = h_fill, h_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row, cnt = 2, 1
        for sale in sales_data:
            depo_str = "TR Depo" if sale.get("source") == "TR_DEPO" else "Özbek (Elips)"
            for item in sale.get("items", []):
                ws.append([cnt, str(sale.get("date", "")), depo_str, str(sale.get("ref_no", "YOK")), str(sale.get("firm", "")),
                           item.get("code", ""), item.get("name", ""), item.get("qty", 1)])
                for c in range(1, len(headers) + 1): ws.cell(row=row, column=c).border = border
                row += 1; cnt += 1
        for idx, w in enumerate([8, 16, 14, 14, 25, 16, 40, 10], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
        wb.save(fn)
        return fn
    except Exception as e:
        logging.error(f"Excel hatası: {e}")
        return None

def generate_pdf_z_report(sales_data, title="Z_Raporu"):
    fn = f"Elips_{title}_{datetime.now().strftime('%d_%m_%Y')}.pdf"
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        doc = SimpleDocTemplate(fn, pagesize=A4, rightMargin=15, leftMargin=15, topMargin=20, bottomMargin=20)
        styles = getSampleStyleSheet()
        t_style = ParagraphStyle('T', parent=styles['Heading1'], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor('#161F30'), alignment=1, spaceAfter=4)
        s_style = ParagraphStyle('S', parent=styles['Normal'], fontName="Helvetica", fontSize=8, textColor=colors.HexColor('#64748B'), alignment=1, spaceAfter=12)
        elements = [
            Paragraph("ELIPS TEKSTIL ELIAR PARCA SATIS RAPORU", t_style),
            Paragraph(f"RESMI DOKUM TARIHI: {datetime.now().strftime('%d.%m.%Y %H:%M')}", s_style)
        ]
        data = [["Sira", "Tarih", "Depo", "Musteri / Tesis", "Urun Kodu", "Urun Adi", "Adet"]]
        cnt, total = 1, 0
        for sale in sales_data:
            depo = "TR" if sale.get("source") == "TR_DEPO" else "Elips"
            firm = str(sale.get("firm", ""))[:22]
            date = str(sale.get("date", "")).split()[0]
            for itm in sale.get("items", []):
                q = int(itm.get("qty", 1))
                total += q
                data.append([str(cnt), date, depo, firm, str(itm.get("code", "")), str(itm.get("name", ""))[:28], str(q)])
                cnt += 1
        data.append(["TOPLAM", "", "", "", "", "Toplam Adet:", str(total)])
        table = Table(data, colWidths=[30, 65, 45, 140, 75, 155, 45], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#161F30')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E2E8F0')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold')
        ]))
        elements.append(table)
        doc.build(elements)
        return fn
    except Exception as e:
        logging.error(f"PDF hatası: {e}")
        return None

def get_client_history_text(firm_query: str):
    sales = load_sales()
    q_norm = normalize_tr(firm_query.strip())
    scored = []
    for s in sales:
        if s.get("status") != "APPROVED": continue
        f_norm = normalize_tr(s.get("firm", ""))
        ratio = difflib.SequenceMatcher(None, q_norm, f_norm).ratio()
        if q_norm in f_norm or ratio > 0.40:
            scored.append((s, ratio))

    if not scored:
        return f"🔍 <code>{html.escape(firm_query)}</code> için geçmiş onaylı kayıt bulunamadı."

    scored.sort(key=lambda x: x[1], reverse=True)
    best_firm = scored[0][0].get("firm")
    firm_records = [s for s, _ in scored if s.get("firm") == best_firm]

    text = f"🏢 <b>{html.escape(best_firm)} — PARÇA GEÇMİŞİ</b>\n═══════════════════════════════\n\n"
    for item in firm_records:
        depo_lbl = " [TR]" if item.get("source") == "TR_DEPO" else " [Elips]"
        for itm in item.get("items", []):
            text += f"• 📅 <code>{item.get('date', '')[:10]}</code> ➔ <b>{itm.get('qty')}x</b> <code>[{itm.get('code')}]</code> {html.escape(itm.get('name', ''))}{depo_lbl}\n"
    text += f"\n═══════════════════════════════\n✅ Toplam: <b>{len(firm_records)} Sevkiyat / Hareket</b>"
    return text

# ----------------- MENÜ DÜZENLEMESİ (is_persistent=False) -----------------
def get_main_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="⚡ Eliar Mini App", web_app=WebAppInfo(url=WEBAPP_URL))],
            [KeyboardButton(text="📝 Eliar'dan Teklif Al"), KeyboardButton(text="📦 Parça Listesi")],
            [KeyboardButton(text="📊 Z-Raporu Al"), KeyboardButton(text="📅 Tarihli Rapor Al")],
            [KeyboardButton(text="⏳ Bekleyen Teklifler"), KeyboardButton(text="✅ Tamamlanan Teklifler")],
            [KeyboardButton(text="📂 Verilen Teklifler"), KeyboardButton(text="🏢 Tesis Geçmişi")],
            [KeyboardButton(text="📖 Kataloğu Gönder")]
        ],
        resize_keyboard=True,
        is_persistent=False,
        input_field_placeholder="İşlem seçin..."
    )

def get_start_inline_keyboard():
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="⚡ Eliar Mini App'i Aç", web_app=WebAppInfo(url=WEBAPP_URL))],
            [InlineKeyboardButton(text="📝 Eliar'dan Teklif Al", callback_data="btn_teklif_hazirla")],
            [InlineKeyboardButton(text="📦 Parça Listesi", callback_data="btn_parca_listesi"), InlineKeyboardButton(text="📊 Z-Raporu Al", callback_data="btn_z_raporu")],
            [InlineKeyboardButton(text="⏳ Bekleyen Teklifler", callback_data="btn_bekleyenler"), InlineKeyboardButton(text="✅ Tamamlanan Teklifler", callback_data="btn_tamamlananlar")],
            [InlineKeyboardButton(text="📂 Verilen Teklifler", callback_data="btn_verilenler"), InlineKeyboardButton(text="🏢 Tesis Geçmişi", callback_data="btn_tesis_gecmisi")],
            [InlineKeyboardButton(text="📖 Kataloğu İndir", callback_data="btn_katalog_indir")]
        ]
    )

async def setup_bot_commands():
    commands = [
        BotCommand(command="start", description="Başlat / Ana Menü"),
        BotCommand(command="teklif", description="Eliar'dan Teklif İste"),
        BotCommand(command="stok", description="Parça Listesi"),
        BotCommand(command="bekleyenler", description="Bekleyen Teklifler"),
        BotCommand(command="tamamlananlar", description="Tamamlanan Teklifler"),
        BotCommand(command="verilenler", description="Verilen Teklifler"),
        BotCommand(command="zraporu", description="Satış Z-Raporu"),
        BotCommand(command="tarihraporu", description="Tarihli Rapor Al"),
        BotCommand(command="gecmis", description="Tesis Geçmişi"),
        BotCommand(command="katalog", description="Katalog PDF İndir")
    ]
    await bot.set_my_commands(commands)
    try:
        await bot.set_chat_menu_button(menu_button=types.MenuButtonCommands())
    except Exception as e:
        logging.warning(f"Menu button: {e}")

# ----------------- ANA MENÜ / START -----------------
@dp.message(CommandStart(), StateFilter("*"))
@dp.message(Command("start"), StateFilter("*"))
@dp.message(F.text.casefold().in_(["/start", "start", "başlat", "baslat", "menü", "menu"]), StateFilter("*"))
async def cmd_start(message: types.Message, state: FSMContext):
    await state.clear()
    user_id = message.chat.id
    sender_name = html.escape(message.from_user.first_name if message.from_user else "Kullanıcı")
    name = "Ömer" if user_id == OMER_CHAT_ID else ("Kamil" if user_id == KAMIL_CHAT_ID else sender_name)
    pdf_info = "✓ Bulundu" if get_catalog_pdf_file() else "✗ Bulunamadı (Katalog.pdf)"
    stocks = load_stocks()

    test_banner = "\n🧪 <b>[DİKKAT: TEST / SİMÜLASYON MODU AKTİF - Kamil'e bildirim gitmez, stoklar değişmez!]</b>\n" if TEST_MODE else ""

    msg_text = (
        f"👋 Selam {name}! <b>Elips Depo</b> Robotu Aktif.{test_banner}\n"
        f"• Katalog PDF: <b>{pdf_info}</b>\n"
        f"• Güncel Depo Kalem: <b>{len(stocks)} Çeşit Parça (Tüm Katalog Aktif!)</b>\n\n"
        f"💡 <i>Menüden işlem seçebilir, teklif hazırlayabilir, parça/etiket fotoğrafı atabilir veya teklif PDF'i yükleyebilirsiniz.</i>"
    )
    await message.answer(msg_text, reply_markup=get_start_inline_keyboard(), parse_mode="HTML")
    await message.answer("👇 <b>Menü Butonları:</b>", reply_markup=get_main_keyboard(), parse_mode="HTML")

# ----------------- INLINE BUTON KÖPRÜLERİ -----------------
@dp.callback_query(F.data == "btn_teklif_hazirla")
async def cb_btn_teklif(callback: types.CallbackQuery, state: FSMContext):
    await callback.answer()
    await start_quote_wizard(callback.message, state)

@dp.callback_query(F.data == "btn_parca_listesi")
async def cb_btn_stok(callback: types.CallbackQuery):
    await callback.answer()
    await show_stocks(callback.message)

@dp.callback_query(F.data == "btn_z_raporu")
async def cb_btn_z_raporu(callback: types.CallbackQuery):
    await callback.answer()
    await trigger_z_report_manual(callback.message)

@dp.callback_query(F.data == "btn_bekleyenler")
async def cb_btn_bekleyenler(callback: types.CallbackQuery):
    await callback.answer()
    await show_pending_quotes(callback.message)

@dp.callback_query(F.data == "btn_tamamlananlar")
async def cb_btn_tamamlananlar(callback: types.CallbackQuery):
    await callback.answer()
    await show_completed_quotes(callback.message)

@dp.callback_query(F.data == "btn_verilenler")
async def cb_btn_verilenler(callback: types.CallbackQuery):
    await callback.answer()
    await show_submitted_quotes(callback.message)

@dp.callback_query(F.data == "btn_tesis_gecmisi")
async def cb_btn_tesis_gecmisi(callback: types.CallbackQuery, state: FSMContext):
    await callback.answer()
    await btn_client_history(callback.message, state)

@dp.callback_query(F.data == "btn_katalog_indir")
async def cb_btn_katalog(callback: types.CallbackQuery):
    await callback.answer()
    await send_catalog_pdf(callback.message)

@dp.message(Command("katalog"))
@dp.message(F.text == "📖 Kataloğu Gönder")
async def send_catalog_pdf(message: types.Message):
    cat_path = get_catalog_pdf_file()
    if cat_path and os.path.exists(cat_path):
        await message.answer_document(FSInputFile(cat_path), caption="📖 <b>Elips Depo - Yedek Parça Kataloğu</b>", parse_mode="HTML")
    else:
        await message.answer("⚠️ Sunucuda <code>Katalog.pdf</code> bulunamadı.", parse_mode="HTML")

# ==============================================================================
#      DÜZELTİLMİŞ FOTOĞRAF / ETİKET TANIMA (GEMINI VISION)
# ==============================================================================
@dp.message(F.photo)
async def handle_photo_recognition(message: types.Message, state: FSMContext):
    wait_msg = await message.answer("🔍 <i>Fotoğraftaki parça, etiket ve model bilgisi yapay zeka ile taranıyor...</i>", parse_mode="HTML")

    tmp_img_path = f"tmp_recog_{int(time.time()*1000)}.jpg"
    try:
        photo = message.photo[-1]
        file_obj = await bot.get_file(photo.file_id)
        await bot.download_file(file_obj.file_path, tmp_img_path)

        with open(tmp_img_path, "rb") as f:
            img_bytes = f.read()

        det_code = ""
        det_name = ""

        if ai_client:
            prompt = (
                "Sen endüstriyel tekstil otomasyonu (Eliar / Elips) yedek parça ve donanım uzmanısın. "
                "Fotoğraftaki parçayı, etiketini, modelini, sipariş kodunu veya tipini (vana, sensör, diyafram, sürücü, kart, silindir, bobin, pompa vb.) incele.\n"
                "KESİNLİKLE YALNIZCA SAF JSON DÖNDÜR:\n"
                "{\n"
                '  "detected_code": "Varsa etiket/parça kodu veya model (örn: 90416, 90567, V1000, CD85N, BES005M, 20031 vb.) yoksa boş bırak",\n'
                '  "detected_name": "Parçanın en doğru Türkçe teknik tanımı",\n'
                '  "category": "MEKATRONIK veya ELEKTRONIK veya KONTROL"\n'
                "}"
            )

            try:
                from PIL import Image
                pil_img = Image.open(BytesIO(img_bytes))
            except Exception:
                pil_img = None

            loop = asyncio.get_event_loop()
            content_payload = [prompt, pil_img] if pil_img else [prompt, genai_types.Part.from_bytes(data=img_bytes, mime_type="image/jpeg")]

            for m_name in ["gemini-2.0-flash", "gemini-1.5-flash"]:
                try:
                    response = await loop.run_in_executor(
                        None,
                        lambda: ai_client.models.generate_content(
                            model=m_name,
                            contents=content_payload
                        )
                    )
                    if response and response.text:
                        raw = response.text.replace("```json", "").replace("```", "").strip()
                        parsed = json.loads(raw)
                        det_code = str(parsed.get("detected_code", "")).strip()
                        det_name = str(parsed.get("detected_name", "")).strip()
                        break
                except Exception as ex:
                    logging.warning(f"Vision {m_name} denemesi: {ex}")

        matched_part = None
        stocks = load_stocks()

        if det_code:
            matched_part = next((s for s in stocks if s["code"].lower() == det_code.lower()), None)
        if not matched_part and det_name:
            matches = find_best_matches(det_name)
            if matches: matched_part = matches[0][0]

        await wait_msg.delete()

        if matched_part:
            cat_img = extract_part_image_from_pdf(matched_part["code"], matched_part["name"])
            if cat_img:
                photo_file = BufferedInputFile(cat_img, filename=f"{matched_part['code']}.png")
                await message.answer_photo(photo_file, caption=f"📸 <b>Katalog Eşleşmesi:</b> <code>[{matched_part['code']}]</code> {matched_part['name']}", parse_mode="HTML")

            stock_stat = f"🟢 Mevcut ({matched_part['qty']} ADT)" if matched_part['qty'] > 0 else "🔴 Stokta Yok (0 ADT)"
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="📝 Bu Parça İçin Eliar'dan Teklif Al", callback_data=f"prefill_q:{matched_part['code']}")],
                [InlineKeyboardButton(text="❌ Kapat", callback_data="cb_close_card")]
            ])

            res_text = (
                f"🎯 <b>GÖRSELDEN PARÇA BAŞARIYLA TANINDI!</b>\n"
                f"═══════════════════════════════\n"
                f"🔍 <b>Tespit:</b> {det_name} {f'[{det_code}]' if det_code else ''}\n"
                f"📦 <b>Depo Eşleşmesi:</b> <code>[{matched_part['code']}]</code> {matched_part['name']}\n"
                f"📊 <b>Depo Durumu:</b> <b>{stock_stat}</b> | Konum: <code>{matched_part['loc']}</code>\n"
                f"═══════════════════════════════\n"
                f"❓ <i>Bu parça için Eliar'dan (Nilay Hanım'dan) fiyat teklifi talebi oluşturulsun mu?</i>"
            )
            await message.answer(res_text, reply_markup=kb, parse_mode="HTML")
        else:
            await message.answer(
                f"🔍 <b>Görsel İncelendi:</b>\n"
                f"• Tespit Edilen: <b>{det_name or 'Endüstriyel Parça'}</b> {f'[{det_code}]' if det_code else ''}\n\n"
                f"ℹ️ <i>Depo kayıtlarında otomatik eşleşmedi. Parça kodunu veya ismini yazarak teklif oluşturabilirsiniz.</i>",
                parse_mode="HTML"
            )

    except Exception as e:
        logging.error(f"Görsel analiz hatası: {e}")
        try: await wait_msg.delete()
        except Exception: pass
        await message.answer("⚠️ Görsel taranırken bir hata oluştu. Parça kodunu manuel yazabilirsiniz.")
    finally:
        if os.path.exists(tmp_img_path):
            try: os.remove(tmp_img_path)
            except Exception: pass

@dp.callback_query(F.data.startswith("prefill_q:"))
async def on_prefill_quote(callback: types.CallbackQuery, state: FSMContext):
    code = callback.data.split(":")[1]
    stocks = load_stocks()
    part = next((s for s in stocks if s["code"] == code), None)
    if not part:
        await callback.answer("Parça bulunamadı.", show_alert=True)
        return

    await state.clear()
    await state.update_data(quote_items=[], prefilled_part=part)
    await state.set_state(QuoteBuilderState.waiting_firm)
    await callback.message.answer(
        f"🏢 <code>[{part['code']}]</code> {part['name']} için teklif <b>hangi müşteri / fabrika</b> için hazırlanacak?\n"
        f"(Örn: <code>HAJJ TEKSTİL</code>, <code>TERRY JAR</code>):",
        parse_mode="HTML"
    )
    await callback.answer()

@dp.callback_query(F.data == "cb_close_card")
async def cb_close(callback: types.CallbackQuery):
    await callback.message.delete()
    await callback.answer()

# ==============================================================================
#      DÜZELTİLMİŞ EŞLEŞTİRME: HEM FİRMA HEM DE PARÇALAR KONTROL EDİLİR!
# ==============================================================================
@dp.message(F.document)
async def handle_document_upload(message: types.Message):
    doc = message.document
    if not doc.file_name.lower().endswith(".pdf"):
        await message.answer("⚠️ Lütfen PDF formatında bir teklif dosyası yükleyin.")
        return

    wait_msg = await message.answer("⏳ <i>PDF analiz ediliyor, parça ve fiyat bilgileri doğrulanıyor...</i>", parse_mode="HTML")

    saved_filename = f"{int(time.time())}_{doc.file_name}"
    save_path = os.path.join(PDF_STORAGE_DIR, saved_filename)
    await bot.download(doc, destination=save_path)

    parsed_data = await parse_quote_pdf_smart(save_path, doc.file_name)
    firm = parsed_data.get("firm", "BİLİNMEYEN FİRMA").upper()
    ref_no = parsed_data.get("ref_no", "BELİRTİLMEDİ")
    date = parsed_data.get("date", datetime.now().strftime("%d.%m.%Y"))
    currency = parsed_data.get("currency", "EUR")
    grand_total = parsed_data.get("grand_total", "-")
    items = parsed_data.get("items", [])
    is_nilay = parsed_data.get("is_nilay", False)

    # 1. HEM FİRMA ADI HEM DE PARÇA EŞLEŞMESİ KONTROLÜ (KESİN DOĞRULAMA)
    quotes = load_quotes()
    matched_quote = None

    for q in quotes:
        if q.get("status") == "PENDING":
            f1 = normalize_tr(q.get("firm", ""))
            f2 = normalize_tr(firm)
            firm_match = (f1 in f2 or f2 in f1 or difflib.SequenceMatcher(None, f1, f2).ratio() > 0.60)

            items_match = False
            for req_item in q.get("items", []):
                req_code = str(req_item.get("code", "")).lower()
                req_name = normalize_tr(req_item.get("name", ""))
                for p_itm in items:
                    p_code = str(p_itm.get("code", "")).lower()
                    p_name = normalize_tr(p_itm.get("name", ""))
                    if (req_code and req_code == p_code) or (req_name and req_name in p_name):
                        items_match = True
                        break

            if firm_match and items_match:
                matched_quote = q
                break

    quote_id = f"TK-{int(time.time())}"
    temp_quote = {
        "id": quote_id,
        "firm": firm,
        "ref_no": ref_no,
        "date": date,
        "currency": currency,
        "grand_total": grand_total,
        "status": "PENDING",
        "items": items,
        "pdf_file": save_path
    }

    await wait_msg.delete()

    # EĞER BOT HAZIRLAMADIYSA VEYA PARÇALAR FARKLIYSA -> KULLANICIYA NET SOR!
    if not matched_quote:
        temp_quote["status"] = "UNMATCHED_PROMPT"
        quotes.append(temp_quote)
        save_quotes(quotes)

        prompt_text = (
            f"⚠️ <b>HARİCİ / BOT DIŞI TEKLİF BELGESİ ALGILANDI</b>\n"
            f"═══════════════════════════════\n"
            f"🏢 <b>Müşteri:</b> <b>{html.escape(firm)}</b>\n"
            f"🔖 <b>Ref No:</b> <code>{ref_no}</code> | 📅 <b>Tarih:</b> {date}\n"
            f"💰 <b>Toplam Tutar:</b> <b>{grand_total} {currency}</b>\n"
            f"📦 <b>Tespit Edilen Kalem:</b> {len(items)} Parça\n"
            f"═══════════════════════════════\n"
            f"ℹ️ <i>Bu teklif daha önce Eliar Depo Botu üzerinden hazırlanmamıştır (veya parçaları eşleşmemektedir).</i>\n\n"
            f"❓ <b>Bu teklifi yine de sisteme girmek ve işlemek ister misiniz?</b>"
        )
        kb_prompt = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Evet, Sisteme Gir ve Analiz Et", callback_data=f"qext_accept:{quote_id}")],
            [InlineKeyboardButton(text="❌ Hayır, İptal Et", callback_data=f"qext_reject:{quote_id}")]
        ])
        await message.answer(prompt_text, reply_markup=kb_prompt, parse_mode="HTML")
        return

    # EĞER HEM FİRMA HEM PARÇALAR EŞLEŞTİYSE -> DOĞRULANMIŞ EKRANA GEÇ
    quotes.append(temp_quote)
    save_quotes(quotes)
    await render_quote_routing_screen(message, temp_quote, is_nilay=is_nilay, matched_quote=matched_quote)

async def render_quote_routing_screen(target_msg, quote_data, is_nilay=False, matched_quote=None):
    quote_id = quote_data["id"]
    firm = quote_data["firm"]
    ref_no = quote_data["ref_no"]
    date = quote_data["date"]
    currency = quote_data["currency"]
    grand_total = quote_data["grand_total"]
    items = quote_data.get("items", [])

    match_lines = []
    if matched_quote:
        match_lines.append("🎯 <b>ELİAR'A (NİLAY HANIM'A) HAZIRLANAN TEKLİF İLE EŞLEŞTİ!</b>")
        match_lines.append(f"<i>(Daha önce hazırlanan {matched_quote.get('firm')} parça talebi doğrulandı)</i>")
    elif is_nilay:
        match_lines.append("🎯 <b>NİLAY DEMİR İMZALI RESMİ TEKLİF DOĞRULANDI!</b>")
    else:
        match_lines.append("📄 <b>SİSTEME ALINAN FİYAT TEKLİFİ</b>")

    stocks = load_stocks()
    stock_analysis_lines = []

    for itm in items:
        code = itm.get("code", "")
        req_qty = int(itm.get("qty", 1))
        p = next((s for s in stocks if s["code"].lower() == code.lower()), None)
        if not p:
            p = next((s for s in stocks if normalize_tr(s["name"]) in normalize_tr(itm.get("name", ""))), None)

        avail_qty = p.get("qty", 0) if p else 0
        itm["depo_available"] = avail_qty

        if avail_qty >= req_qty:
            stock_analysis_lines.append(f"• <code>[{code}]</code> {html.escape(itm.get('name', ''))} ({req_qty} Adet) ➔ 🟢 <b>Depoda Var ({avail_qty} ADT)</b>")
        elif avail_qty > 0:
            stock_analysis_lines.append(f"• <code>[{code}]</code> {html.escape(itm.get('name', ''))} ({req_qty} Adet) ➔ 🟡 <b>Kısmi Stok ({avail_qty} ADT Var)</b>")
        else:
            stock_analysis_lines.append(f"• <code>[{code}]</code> {html.escape(itm.get('name', ''))} ({req_qty} Adet) ➔ 🔴 <b>Depoda YOK (0 ADT)</b>")

    lines = match_lines + [
        "═══════════════════════════════",
        f"🏢 <b>Müşteri:</b> <b>{html.escape(firm)}</b>",
        f"🔖 <b>Ref No:</b> <code>{ref_no}</code> | 📅 <b>Tarih:</b> {date}",
        f"💰 <b>Toplam Tutar:</b> <b>{grand_total} {currency}</b>\n",
        "<b>📦 ÖZBEKİSTAN (ELİPS) DEPO ENVANTER ANALİZİ:</b>"
    ]
    lines.extend(stock_analysis_lines if stock_analysis_lines else ["⚠️ <i>Kalem tablosu tespit edilemedi.</i>"])
    lines.append("\n═══════════════════════════════")
    lines.append("❓ <b>Bu parçalar nereden tedarik edilecek / sevk edilecek?</b>")

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🇺🇿 Özbek (Elips) Depodan Karşıla", callback_data=f"qroute:{quote_id}:ELIPS")],
        [InlineKeyboardButton(text="🇹🇷 Eliar TR'den Gelecek", callback_data=f"qroute:{quote_id}:TR")],
        [InlineKeyboardButton(text="🔄 Kısmi (Elimizdekiler Özbek, Kalan TR)", callback_data=f"qroute:{quote_id}:PARTIAL")],
        [InlineKeyboardButton(text="❌ Teklifi Kabul Etme / İptal", callback_data=f"qstat:{quote_id}:REJECTED")]
    ])

    if hasattr(target_msg, "answer"):
        await target_msg.answer("\n".join(lines), reply_markup=kb, parse_mode="HTML")
    else:
        await target_msg.edit_text("\n".join(lines), reply_markup=kb, parse_mode="HTML")

@dp.callback_query(F.data.startswith("qext_accept:"))
async def on_quote_external_accept(callback: types.CallbackQuery):
    quote_id = callback.data.split(":")[1]
    quotes = load_quotes()
    q = next((x for x in quotes if x["id"] == quote_id), None)
    if not q:
        await callback.answer("Teklif bulunamadı.", show_alert=True)
        return
    q["status"] = "PENDING"
    save_quotes(quotes)
    await callback.answer("Teklif sisteme alındı.")
    await render_quote_routing_screen(callback.message, q)

@dp.callback_query(F.data.startswith("qext_reject:"))
async def on_quote_external_reject(callback: types.CallbackQuery):
    quote_id = callback.data.split(":")[1]
    quotes = load_quotes()
    quotes = [x for x in quotes if x["id"] != quote_id]
    save_quotes(quotes)
    await callback.message.edit_text("❌ <b>Harici teklif sisteme kaydedilmedi, işlem iptal edildi.</b>", parse_mode="HTML")
    await callback.answer("İptal edildi.")

# ----------------- SEVKİYAT MENŞEİ (ÖZBEK / TR) SEÇİMİ -----------------
@dp.callback_query(F.data.startswith("qroute:"))
async def on_quote_route_selected(callback: types.CallbackQuery):
    _, quote_id, route_choice = callback.data.split(":")
    quotes = load_quotes()
    q = next((x for x in quotes if x["id"] == quote_id), None)
    if not q:
        await callback.answer("Teklif kaydı bulunamadı.", show_alert=True)
        return

    stocks = load_stocks()
    items = q.get("items", [])

    if route_choice == "TR":
        q["status"] = "APPROVED"
        q["source"] = "TR_DEPO"
        q["completed_at"] = datetime.now().isoformat()
        save_quotes(quotes)

        sales = load_sales()
        sales.append({
            "id": f"SLS-{int(time.time())}",
            "firm": q["firm"],
            "source": "TR_DEPO",
            "ref_no": q.get("ref_no", "TR-SEVKIYAT"),
            "items": items,
            "status": "APPROVED",
            "date": datetime.now().strftime("%d.%m.%Y %H:%M"),
            "timestamp": datetime.now().isoformat()
        })
        save_sales(sales)

        msg = (
            f"🇹🇷 <b>TEKLİF ONAYLANDI (ELİAR TÜRKİYE'DEN GELECEK)!</b>\n"
            f"═══════════════════════════════\n"
            f"🏢 Firma: <b>{html.escape(q['firm'])}</b>\n"
            f"🔖 Ref No: <code>{q.get('ref_no', 'YOK')}</code>\n"
            f"📦 Parçalar Türkiye merkez depodan sevk edilecektir.\n"
            f"✅ <i>Özbekistan Elips depo stoklarına dokunulmadı.</i>\n\n"
            f"📌 <i>Teklif 'Tamamlanan Teklifler' menüsüne işlendi.</i>"
        )
        await notify_both(msg)
        await callback.message.edit_text(msg, parse_mode="HTML")
        await callback.answer("Türkiye sevkiyatı olarak onaylandı.")
        return

    missing_items = []
    available_items = []

    for itm in items:
        code = itm.get("code", "")
        req_qty = int(itm.get("qty", 1))
        p = next((s for s in stocks if s["code"].lower() == code.lower()), None)
        avail = p.get("qty", 0) if p else 0

        if avail >= req_qty:
            available_items.append((code, itm.get("name", ""), req_qty, avail))
        elif avail > 0:
            available_items.append((code, itm.get("name", ""), avail, avail))
            missing_items.append((code, itm.get("name", ""), req_qty - avail))
        else:
            missing_items.append((code, itm.get("name", ""), req_qty))

    q["route_mode"] = route_choice
    save_quotes(quotes)

    warn_text = ""
    if missing_items:
        warn_text = "\n\n⚠️ <b>DİKKAT: Bazı parçalar Özbekistan Elips depomuzda YOK:</b>\n"
        for c, n, mq in missing_items:
            warn_text += f"• <code>[{c}]</code> {n} (<b>{mq} Adet Eksik</b> ➔ TR'den gelmeli)\n"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Evet, Elips Stoğundan Düş", callback_data=f"qapplystock:{quote_id}:yes")],
        [InlineKeyboardButton(text="❌ Hayır, Sadece Durumu Onayla", callback_data=f"qapplystock:{quote_id}:no")]
    ])

    await callback.message.edit_text(
        f"🇺🇿 <b>ÖZBEKİSTAN (ELİPS) DEPO ÇIKIŞ ONAYI</b>\n"
        f"🏢 Firma: <b>{html.escape(q['firm'])}</b>"
        f"{warn_text}\n"
        f"Mevcut olan parçalar Elips Depo stoğundan düşülsün mü?",
        reply_markup=kb,
        parse_mode="HTML"
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("qapplystock:"))
async def on_quote_apply_stock(callback: types.CallbackQuery):
    _, quote_id, choice = callback.data.split(":")
    quotes = load_quotes()
    q = next((x for x in quotes if x["id"] == quote_id), None)
    if not q:
        await callback.answer("Teklif bulunamadı.", show_alert=True)
        return

    q["status"] = "APPROVED"
    q["source"] = "ELIPS_DEPO"
    q["completed_at"] = datetime.now().isoformat()
    save_quotes(quotes)

    if choice == "yes" and q.get("items"):
        stocks = load_stocks()
        sales = load_sales()
        deducted_report = []
        not_found_report = []

        for item in q["items"]:
            code = item.get("code", "")
            req_qty = int(item.get("qty", 1))

            p = next((s for s in stocks if s["code"].lower() == code.lower()), None)
            if not p:
                p = next((s for s in stocks if normalize_tr(s["name"]) in normalize_tr(item.get("name", ""))), None)

            if p and p.get("qty", 0) > 0:
                actual_deduct = min(p["qty"], req_qty)
                p["qty"] -= actual_deduct
                deducted_report.append(f"• <code>[{p['code']}]</code> {p['name']}: <b>-{actual_deduct} Adet</b> (Kalan: {p['qty']})")
                if actual_deduct < req_qty:
                    not_found_report.append(f"• <code>[{p['code']}]</code> {p['name']}: {req_qty - actual_deduct} Adet depoda yok (TR'den gelmeli)")
            else:
                not_found_report.append(f"• <code>[{code}]</code> {item.get('name', '')}: Depoda 0 adet var (TR'den gelmeli)")

        save_stocks(stocks)

        sales.append({
            "id": f"SLS-{int(time.time())}",
            "firm": q["firm"],
            "source": "ELIPS_DEPO",
            "ref_no": q.get("ref_no", "TEKLIF-ONAY"),
            "items": q["items"],
            "status": "APPROVED",
            "date": datetime.now().strftime("%d.%m.%Y %H:%M"),
            "timestamp": datetime.now().isoformat()
        })
        save_sales(sales)

        broadcast = (
            f"📦 <b>ELİPS DEPO - STOK ÇIKIŞI GERÇEKLEŞTİ!</b>\n"
            f"═══════════════════════════════\n"
            f"🏢 Firma: <b>{html.escape(q['firm'])}</b>\n"
            f"🔖 Ref No: <code>{q.get('ref_no', 'YOK')}</code>\n\n"
            f"<b>✅ Düşülen Parçalar:</b>\n" + ("\n".join(deducted_report) if deducted_report else "Yok")
        )
        if not_found_report:
            broadcast += "\n\n<b>⚠️ Depoda Olmayan / TR'den İstenecekler:</b>\n" + "\n".join(not_found_report)

        broadcast += "\n\n📌 <i>Teklif 'Tamamlanan Teklifler' menüsüne taşındı.</i>"
        await notify_both(broadcast)
        await callback.message.edit_text(broadcast, parse_mode="HTML")
    else:
        await callback.message.edit_text("✅ <b>Teklif onaylandı (Depo stokları değiştirilmedi).</b>", parse_mode="HTML")
    await callback.answer("İşlem tamamlandı.")

@dp.callback_query(F.data.startswith("qstat:"))
async def on_quote_status_cancel(callback: types.CallbackQuery):
    _, quote_id, new_status = callback.data.split(":")
    quotes = load_quotes()
    q = next((x for x in quotes if x["id"] == quote_id), None)
    if q:
        q["status"] = new_status
        save_quotes(quotes)
    await callback.message.edit_text("🔴 <b>Teklif reddedildi / iptal edildi.</b>", parse_mode="HTML")
    await callback.answer("İptal edildi.")

# ==============================================================================
#            TAMAMLANMIŞ TEKLİFLER EKRANI & ARAMA / FİLTRELEME
# ==============================================================================
def render_completed_quotes_list(quotes_list):
    buttons = [
        [
            InlineKeyboardButton(text="🔍 Firma ile Ara", callback_data="cq_search_firm"),
            InlineKeyboardButton(text="📅 Tarihe Göre Filtrele", callback_data="cq_filter_date")
        ]
    ]
    for q in reversed(quotes_list[-15:]):
        depo_lbl = " [TR]" if q.get("source") == "TR_DEPO" else " [Elips]"
        firm_title = f"✅ {q['firm']}{depo_lbl} ({q.get('date', '')[:10]})"
        buttons.append([InlineKeyboardButton(text=firm_title, callback_data=f"qview:{q['id']}")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

@dp.message(Command("tamamlananlar"))
@dp.message(F.text == "✅ Tamamlanan Teklifler")
async def show_completed_quotes(message: types.Message):
    quotes = load_quotes()
    approved = [q for q in quotes if q.get("status") == "APPROVED"]
    if not approved:
        await message.answer("ℹ️ Henüz onaylanmış / tamamlanmış bir teklif bulunmamaktadır.", parse_mode="HTML")
        return

    kb = render_completed_quotes_list(approved)
    await message.answer(f"✅ <b>TAMAMLANMIŞ / ONAYLANAN TEKLİFLER ({len(approved)} Adet):</b>\nDetayını veya PDF'ini görmek için seçin:", reply_markup=kb, parse_mode="HTML")

@dp.callback_query(F.data == "cq_search_firm")
async def on_cq_search_firm_prompt(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(CompletedQuoteSearchState.waiting_firm_query)
    await callback.message.answer("🔍 Tamamlanan tekliflerde aramak istediğiniz <b>Firma Adını</b> yazın:\n(Örn: <code>HAJJ</code>, <code>TERRY JAR</code>)", parse_mode="HTML")
    await callback.answer()

@dp.message(CompletedQuoteSearchState.waiting_firm_query, ~F.text.startswith("/"))
async def process_cq_firm_search(message: types.Message, state: FSMContext):
    await state.clear()
    query = normalize_tr(message.text.strip())
    quotes = load_quotes()
    matched = [q for q in quotes if q.get("status") == "APPROVED" and query in normalize_tr(q.get("firm", ""))]

    if not matched:
        await message.answer(f"🔍 <code>{html.escape(message.text)}</code> firmasına ait tamamlanmış teklif bulunamadı.", parse_mode="HTML")
        return

    kb = render_completed_quotes_list(matched)
    await message.answer(f"🔍 <b>'{html.escape(message.text)}' Arama Sonuçları ({len(matched)} Teklif):</b>", reply_markup=kb, parse_mode="HTML")

@dp.callback_query(F.data == "cq_filter_date")
async def on_cq_filter_date_menu(callback: types.CallbackQuery):
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="🟢 Bugünün Teklifleri", callback_data="cq_apply_date:today"),
            InlineKeyboardButton(text="📅 Bu Hafta (7 Gün)", callback_data="cq_apply_date:week")
        ],
        [
            InlineKeyboardButton(text="🗓️ Bu Ay", callback_data="cq_apply_date:month"),
            InlineKeyboardButton(text="📦 Tüm Tamamlananlar", callback_data="cq_apply_date:all")
        ]
    ])
    await callback.message.edit_text("📅 <b>Hangi dönemin tamamlanan tekliflerini görmek istersiniz?</b>", reply_markup=kb, parse_mode="HTML")
    await callback.answer()

@dp.callback_query(F.data.startswith("cq_apply_date:"))
async def on_cq_apply_date_filter(callback: types.CallbackQuery):
    scope = callback.data.split(":")[1]
    now = datetime.now()

    if scope == "today":
        start_dt = datetime(now.year, now.month, now.day)
    elif scope == "week":
        start_dt = now - timedelta(days=7)
    elif scope == "month":
        start_dt = datetime(now.year, now.month, 1)
    else:
        start_dt = datetime(2020, 1, 1)

    quotes = load_quotes()
    filtered = []
    for q in quotes:
        if q.get("status") != "APPROVED": continue
        d = parse_date_safe(q.get("completed_at") or q.get("date"))
        if d and d >= start_dt:
            filtered.append(q)

    if not filtered:
        await callback.answer("Bu dönemde tamamlanan teklif bulunamadı.", show_alert=True)
        return

    kb = render_completed_quotes_list(filtered)
    await callback.message.edit_text(f"📅 <b>Filtrelenen Tamamlanmış Teklifler ({len(filtered)} Adet):</b>", reply_markup=kb, parse_mode="HTML")
    await callback.answer()

# ----------------- BEKLEYEN VE VERİLEN TEKLİFLER -----------------
@dp.message(Command("bekleyenler"))
@dp.message(F.text == "⏳ Bekleyen Teklifler")
async def show_pending_quotes(message: types.Message):
    quotes = load_quotes()
    pending = [q for q in quotes if q.get("status") == "PENDING"]
    if not pending:
        await message.answer("ℹ️ Şu anda bekleyen bir teklif bulunmamaktadır.", parse_mode="HTML")
        return

    buttons = []
    for q in pending[-10:]:
        buttons.append([InlineKeyboardButton(text=f"⏳ {q['firm']} ({q.get('date', '')[:10]})", callback_data=f"qview:{q['id']}")])
    await message.answer("⏳ <b>BEKLEYEN TEKLİFLER:</b>\nİncelemek istediğinize tıklayın:", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons), parse_mode="HTML")

@dp.message(Command("verilenler"))
@dp.message(F.text == "📂 Verilen Teklifler")
async def show_submitted_quotes(message: types.Message):
    quotes = load_quotes()
    if not quotes:
        await message.answer("ℹ️ Henüz kayıtlı bir teklif bulunmamaktadır.", parse_mode="HTML")
        return

    buttons = []
    for q in reversed(quotes[-12:]):
        st_icon = "🟢" if q.get("status") == "APPROVED" else ("🔴" if q.get("status") == "REJECTED" else "⏳")
        depo_lbl = " [TR]" if q.get("source") == "TR_DEPO" else (" [Elips]" if q.get("source") == "ELIPS_DEPO" else "")
        buttons.append([InlineKeyboardButton(text=f"{st_icon} {q['firm']}{depo_lbl} ({q.get('date', '')[:10]})", callback_data=f"qview:{q['id']}")])
    await message.answer("📂 <b>TÜM TEKLİFLER LİSTESİ:</b>", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons), parse_mode="HTML")

@dp.callback_query(F.data.startswith("qview:"))
async def on_quote_view_details(callback: types.CallbackQuery):
    quote_id = callback.data.split(":")[1]
    quotes = load_quotes()
    q = next((x for x in quotes if x["id"] == quote_id), None)
    if not q:
        await callback.answer("Teklif bulunamadı.", show_alert=True)
        return

    st_map = {"PENDING": "⏳ Beklemede", "APPROVED": "✅ Tamamlandı / Onaylandı", "REJECTED": "❌ Reddedildi"}
    cur_st = st_map.get(q.get("status"), "Belirsiz")
    cur = q.get("currency", "EUR")
    depo_str = "🇹🇷 Türkiye (TR) Depo" if q.get("source") == "TR_DEPO" else ("🇺🇿 Özbek (Elips) Depo" if q.get("source") == "ELIPS_DEPO" else "Belirlenmedi")

    lines = [
        f"📄 <b>Teklif Detayı</b>",
        f"🏢 Firma: <b>{html.escape(q['firm'])}</b>",
        f"🔖 Ref No: <code>{q.get('ref_no', 'YOK')}</code>",
        f"📅 Tarih: <code>{q.get('date')}</code>",
        f"📊 Durum: <b>{cur_st}</b>",
        f"📍 Kaynak: <b>{depo_str}</b>",
        f"💰 Toplam: <b>{q.get('grand_total', '-')} {cur}</b>\n",
        "<b>📦 Parçalar:</b>"
    ]
    for idx, itm in enumerate(q.get("items", []), 1):
        lines.append(f"{idx}. <code>[{itm.get('code')}]</code> {itm.get('name')} — <b>{itm.get('qty')} Adet</b> ({itm.get('unit_price', '-')} {cur})")

    if q.get("pdf_file") and os.path.exists(q["pdf_file"]):
        await callback.message.answer_document(FSInputFile(q["pdf_file"]), caption="\n".join(lines), parse_mode="HTML")
    else:
        await callback.message.answer("\n".join(lines), parse_mode="HTML")
    await callback.answer()

# ----------------- ELİAR'DAN TEKLİF ALMA SİHİRBAZI -----------------
@dp.message(Command("teklif", ignore_mention=True))
@dp.message(F.text == "📝 Eliar'dan Teklif Al")
async def start_quote_wizard(message: types.Message, state: FSMContext):
    data = await state.get_data()
    prefilled_part = data.get("prefilled_part")

    await state.clear()
    await state.update_data(quote_items=[])
    if prefilled_part:
        await state.update_data(selected_part=prefilled_part)

    await state.set_state(QuoteBuilderState.waiting_firm)
    await message.answer("🏢 <b>Teklif hangi müşteri / fabrika için hazırlanacak?</b>\n(Örnek: <code>HAJJ TEKSTİL</code>, <code>TERRY JAR</code>)", parse_mode="HTML")

@dp.message(QuoteBuilderState.waiting_firm, ~F.text.startswith("/"))
async def process_quote_firm(message: types.Message, state: FSMContext):
    firm = message.text.strip().upper()
    await state.update_data(firm_name=firm)
    data = await state.get_data()

    if data.get("selected_part"):
        part = data["selected_part"]
        await state.set_state(QuoteBuilderState.waiting_qty)
        await message.answer(f"🎯 <b>Seçilen Parça:</b> <code>[{part['code']}]</code> {html.escape(part['name'])}\n\n🔢 Kaç <b>adet</b> teklif istenecek?", parse_mode="HTML")
        return

    await state.set_state(QuoteBuilderState.waiting_item_query)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📖 Parça Kataloğunu İncele (PDF)", callback_data="qb_get_catalog")]
    ])
    await message.answer(
        f"🔍 <b>Hangi parça için teklif istenecek?</b>\n"
        f"• Varsa <b>Sipariş Kodunu</b> yazın (Örn: <code>91661</code>, <code>90567</code>, <code>40647</code>)\n"
        f"• Veya <b>Ürün İsmini</b> yazın (Örn: <code>diyafram</code>, <code>pin silindir</code>):",
        reply_markup=kb,
        parse_mode="HTML"
    )

@dp.callback_query(F.data == "qb_get_catalog")
async def cb_send_catalog_in_quote(callback: types.CallbackQuery):
    cat_path = get_catalog_pdf_file()
    if cat_path and os.path.exists(cat_path):
        await callback.message.answer_document(FSInputFile(cat_path), caption="📖 <b>Yedek Parça Kataloğu</b>", parse_mode="HTML")
    else:
        await callback.answer("Katalog dosyası bulunamadı.", show_alert=True)
    await callback.answer()

def render_search_results_page(matches, page_idx, query):
    page_size = 5
    total_pages = (len(matches) + page_size - 1) // page_size
    page_idx = max(0, min(page_idx, total_pages - 1))
    current_items = matches[page_idx * page_size : (page_idx + 1) * page_size]

    buttons = [[InlineKeyboardButton(text=f"👉 {p['code']} - {p['name'][:30]}", callback_data=f"qbsel:{p['code']}")] for p, _ in current_items]

    nav_row = []
    if page_idx > 0: nav_row.append(InlineKeyboardButton(text="⬅️ Önceki", callback_data=f"qbpage:{page_idx-1}"))
    if page_idx < total_pages - 1: nav_row.append(InlineKeyboardButton(text="Sonraki ➡️", callback_data=f"qbpage:{page_idx+1}"))
    if nav_row: buttons.append(nav_row)

    buttons.append([
        InlineKeyboardButton(text="🔍 Tekrar Ara", callback_data="qb_retry_search"),
        InlineKeyboardButton(text="📖 Kataloğu Aç", callback_data="qb_get_catalog")
    ])

    text = f"🔍 <b>Arama:</b> <code>{html.escape(query)}</code> (Sayfa: {page_idx + 1}/{total_pages})\nLütfen listeden parçayı seçin:"
    return text, InlineKeyboardMarkup(inline_keyboard=buttons)

@dp.message(QuoteBuilderState.waiting_item_query, ~F.text.startswith("/"))
async def process_quote_search(message: types.Message, state: FSMContext):
    query = message.text.strip()
    matches = find_best_matches(query)

    if not matches:
        kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="📖 Kataloğu Gönder", callback_data="qb_get_catalog")]])
        await message.answer(f"⚠️ <code>{html.escape(query)}</code> ile eşleşen parça bulunamadı.\nKataloğu incelemek ister misiniz?", reply_markup=kb, parse_mode="HTML")
        return

    top_part, _ = matches[0]
    if query.lower() == top_part['code'].lower():
        await state.update_data(selected_part=top_part)
        await state.set_state(QuoteBuilderState.waiting_qty)

        img_bytes = extract_part_image_from_pdf(top_part["code"], top_part.get("name", ""))
        if img_bytes:
            photo = BufferedInputFile(img_bytes, filename=f"{top_part['code']}.png")
            await message.answer_photo(photo, caption=f"📸 <b>Katalog Görseli:</b> <code>[{top_part['code']}]</code> — {html.escape(top_part['name'])}", parse_mode="HTML")

        stock_text = f"Depo: {top_part.get('qty', 0)} ADT" if top_part.get('qty', 0) > 0 else "Stokta Yok (0 ADT)"
        await message.answer(f"🎯 <b>Seçilen:</b> <code>[{top_part['code']}]</code> {html.escape(top_part['name'])}\n📦 <b>{stock_text}</b>\n\n🔢 Kaç <b>adet</b> teklif istenecek?", parse_mode="HTML")
        return

    await state.update_data(search_matches=matches, search_query=query)
    text, kb = render_search_results_page(matches, 0, query)
    await message.answer(text, reply_markup=kb, parse_mode="HTML")

@dp.callback_query(F.data.startswith("qbpage:"))
async def on_quote_page_navigate(callback: types.CallbackQuery, state: FSMContext):
    page_idx = int(callback.data.split(":")[1])
    data = await state.get_data()
    text, kb = render_search_results_page(data.get("search_matches", []), page_idx, data.get("search_query", ""))
    await callback.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
    await callback.answer()

@dp.callback_query(F.data == "qb_retry_search")
async def on_quote_retry(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(QuoteBuilderState.waiting_item_query)
    await callback.message.edit_text("🔍 Lütfen parça ismini veya kodunu tekrar yazın:")
    await callback.answer()

@dp.callback_query(F.data.startswith("qbsel:"))
async def on_quote_part_confirmed(callback: types.CallbackQuery, state: FSMContext):
    code = callback.data.split(":")[1]
    stocks = load_stocks()
    part = next((p for p in stocks if p["code"] == code), None)
    if not part:
        await callback.answer("Parça bulunamadı.", show_alert=True)
        return

    await state.update_data(selected_part=part)
    await state.set_state(QuoteBuilderState.waiting_qty)

    img_bytes = extract_part_image_from_pdf(part["code"], part.get("name", ""))
    if img_bytes:
        photo = BufferedInputFile(img_bytes, filename=f"{part['code']}.png")
        await callback.message.answer_photo(photo, caption=f"📸 <b>Katalog Görseli:</b> <code>[{part['code']}]</code> — {html.escape(part['name'])}", parse_mode="HTML")

    stock_text = f"Depo: {part.get('qty', 0)} ADT" if part.get('qty', 0) > 0 else "Stokta Yok (0 ADT)"
    await callback.message.answer(f"🎯 <b>Seçilen:</b> <code>[{part['code']}]</code> {html.escape(part['name'])}\n📦 <b>{stock_text}</b>\n\n🔢 Kaç <b>adet</b> teklif istenecek?", parse_mode="HTML")
    await callback.answer()

@dp.message(QuoteBuilderState.waiting_qty, ~F.text.startswith("/"))
async def process_quote_quantity(message: types.Message, state: FSMContext):
    try:
        qty = int(message.text.strip())
        if qty <= 0: raise ValueError
    except Exception:
        await message.answer("⚠️ Lütfen geçerli bir sayı girin (Örn: 1, 2):", parse_mode="HTML")
        return

    data = await state.get_data()
    part = data.get("selected_part")
    items = data.get("quote_items", [])
    items.append({"code": part["code"], "name": part["name"], "qty": qty})
    await state.update_data(quote_items=items)

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Başka Parça Ekle", callback_data="qb_more")],
        [InlineKeyboardButton(text="🏁 Teklifi Bitir ve Metni Al", callback_data="qb_done")]
    ])
    await state.set_state(QuoteBuilderState.confirm_more)
    await message.answer(f"✅ <b>{qty} Adet</b> <code>[{part['code']}]</code> eklendi!\nNe yapmak istersiniz?", reply_markup=kb, parse_mode="HTML")

@dp.callback_query(F.data == "qb_more")
async def on_quote_add_another(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(QuoteBuilderState.waiting_item_query)
    await callback.message.edit_text("🔍 Yeni parçanın <b>kodunu</b> veya <b>ismini</b> yazın:", parse_mode="HTML")
    await callback.answer()

@dp.callback_query(F.data == "qb_done")
async def on_quote_completed(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    firm = html.escape(data.get("firm_name", "MÜŞTERİ"))
    items = data.get("quote_items", [])
    if not items:
        await callback.message.edit_text("⚠️ Teklifte ürün yok.")
        await state.clear()
        return

    # METİNDE NİLAY HANIM HİTABI KORUNDU
    formatted_quote = (
        f"Merhaba Nilay hanim\n\n"
        f"{firm} için aşağıdaki listedeki parçalar için fiyat teklifi verebilir misiniz?\n\n"
    )
    for itm in items:
        formatted_quote += f"• {itm['code']} - {itm['name']} - {itm['qty']} Adet\n"
    formatted_quote += "\nİyi çalışmalar."

    quotes = load_quotes()
    quote_id = f"TK-{int(time.time())}"
    quotes.append({
        "id": quote_id,
        "firm": firm,
        "date": datetime.now().strftime("%d.%m.%Y %H:%M"),
        "status": "PENDING",
        "items": items,
        "pdf_file": None,
        "created_via_bot": True
    })
    save_quotes(quotes)

    text = (
        f"📋 <b>KOPYALAMAYA HAZIR ELİAR (NİLAY HANIM) TEKLİF METNİ:</b>\n"
        f"═══════════════════════════════\n\n"
        f"<code>{html.escape(formatted_quote)}</code>\n\n"
        f"═══════════════════════════════\n"
        f"💡 <i>Metnin üstüne tek tık yaparak kopyalayabilir, Nilay Hanım'a WhatsApp/Mail ile atabilirsiniz.</i>"
    )
    await callback.message.edit_text(text, parse_mode="HTML")

    notify_text = (
        f"📥 <b>YENİ TEKLİF TALEBİ OLUŞTURULDU!</b>\n\n"
        f"🏢 <b>Firma:</b> {firm}\n"
        f"📦 <b>Ürünler:</b>\n" +
        "\n".join([f"• [{i['code']}] {i['name']} ({i['qty']} Adet)" for i in items])
    )
    await notify_both(notify_text)
    await state.clear()
    await callback.answer("Teklif kaydedildi!")

# ----------------- PARÇA LİSTESİ (170+ KALEM KATEGORİLİ VE SAYFALI) -----------------
@dp.message(Command("stok", ignore_mention=True))
@dp.message(F.text == "📦 Parça Listesi")
async def show_stocks(message: types.Message):
    stocks = load_stocks()
    available = [s for s in stocks if s.get("qty", 0) > 0]
    out_of_stock = [s for s in stocks if s.get("qty", 0) == 0]

    text = (
        f"📦 <b>ELİAR DEPO TAM ENVANTER LİSTESİ</b>\n"
        f"📊 Toplam Kalem: <b>{len(stocks)} Çeşit</b> | Stokta: <b>{len(available)}</b> | Tükendi: <b>{len(out_of_stock)}</b>\n"
        f"═══════════════════════════════\n\n"
    )

    for s in available[:15]:
        text += f"• <code>[{s['code']}]</code> {html.escape(s['name'][:35])}\n  └ Stok: <b>{s['qty']} ADT</b> | {s['loc']}\n"
    
    if len(available) > 15:
        text += f"\n<i>...ve {len(available)-15} kalem ürün daha depoda kayıtlıdır. (Tüm 170+ kalemi Mini App'ten canlı yönetebilir veya arama yapabilirsiniz).</i>"

    await message.answer(text, parse_mode="HTML")

# ----------------- Z-RAPORU MENÜSÜ & İŞLEMLERİ -----------------
@dp.message(Command("zraporu", ignore_mention=True))
@dp.message(F.text == "📊 Z-Raporu Al")
async def trigger_z_report_manual(message: types.Message):
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="🟢 Bugünün Z-Raporu", callback_data="date_filter:today"),
            InlineKeyboardButton(text="📅 Bu Hafta (Son 7 Gün)", callback_data="date_filter:7days")
        ],
        [
            InlineKeyboardButton(text="🗓️ Bu Ayın Raporu", callback_data="date_filter:this_month"),
            InlineKeyboardButton(text="📦 Tüm Satışlar & Kapanış", callback_data="date_filter:all")
        ],
        [
            InlineKeyboardButton(text="✍️ Özel Tarih Aralığı Gir", callback_data="date_filter:custom")
        ]
    ])
    await message.answer("📊 <b>Hangi dönemin Satış Z-Raporunu almak istiyorsunuz?</b>", reply_markup=kb, parse_mode="HTML")

@dp.message(Command("tarihraporu", ignore_mention=True))
@dp.message(F.text == "📅 Tarihli Rapor Al")
async def trigger_custom_date_report(message: types.Message):
    await trigger_z_report_manual(message)

@dp.callback_query(F.data.startswith("date_filter:"))
async def process_date_filter_btn(callback: types.CallbackQuery, state: FSMContext):
    mode = callback.data.split(":")[1]
    now = datetime.now()

    if mode == "custom":
        await state.set_state(CustomReportState.waiting_date_range)
        await callback.message.answer("✍️ Format: <code>GG.AA.YYYY - GG.AA.YYYY</code>\n(Örnek: <code>01.10.2026 - 30.10.2026</code>)", parse_mode="HTML")
        await callback.answer()
        return

    if mode == "today":
        s_dt = datetime(now.year, now.month, now.day, 0, 0, 0)
        e_dt = datetime(now.year, now.month, now.day, 23, 59, 59)
        title_label, file_label = "GÜNLÜK Z-RAPORU (BUGÜN)", "Gunluk_Z_Raporu"
    elif mode == "7days":
        s_dt = now - timedelta(days=7)
        e_dt = now
        title_label, file_label = "HAFTALIK Z-RAPORU (SON 7 GÜN)", "Haftalik_Z_Raporu"
    elif mode == "this_month":
        s_dt = datetime(now.year, now.month, 1, 0, 0, 0)
        e_dt = now
        title_label, file_label = "BU AYIN Z-RAPORU", "Aylik_Z_Raporu"
    else:  # all
        s_dt = datetime(2020, 1, 1)
        e_dt = datetime(2035, 1, 1)
        title_label, file_label = "GENEL SATIŞ Z-RAPORU (TÜMÜ)", "Genel_Z_Raporu"

    sales = filter_sales_by_range(s_dt, e_dt)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Excel İndir (.xlsx)", callback_data="dl_rep:excel:custom"),
         InlineKeyboardButton(text="📑 PDF İndir (.pdf)", callback_data="dl_rep:pdf:custom")],
        [InlineKeyboardButton(text="📦 İkisini de İndir (Excel + PDF)", callback_data="dl_rep:both:custom")]
    ])
    await state.update_data(last_filtered_sales=sales, report_title=file_label)
    await callback.message.edit_text(build_report_text(sales, title_label), reply_markup=kb, parse_mode="HTML")
    await callback.answer()

@dp.message(CustomReportState.waiting_date_range, ~F.text.startswith("/"))
async def process_custom_date_range(message: types.Message, state: FSMContext):
    try:
        s_str, e_str = [p.strip() for p in message.text.split("-")]
        s_dt = datetime.strptime(s_str, "%d.%m.%Y")
        e_dt = datetime.strptime(e_str, "%d.%m.%Y").replace(hour=23, minute=59, second=59)
        sales = filter_sales_by_range(s_dt, e_dt)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📊 Excel İndir", callback_data="dl_rep:excel:custom"),
             InlineKeyboardButton(text="📑 PDF İndir", callback_data="dl_rep:pdf:custom")],
            [InlineKeyboardButton(text="📦 İkisi Birden", callback_data="dl_rep:both:custom")]
        ])
        await state.update_data(last_filtered_sales=sales, report_title=f"Rapor_{s_str}_{e_str}")
        await message.answer(build_report_text(sales, f"ÖZEL Z-RAPORU ({s_str} - {e_str})"), reply_markup=kb, parse_mode="HTML")
    except Exception:
        await message.answer("⚠️ Format hatalı! Örnek format: <code>01.10.2026 - 30.10.2026</code>", parse_mode="HTML")

@dp.callback_query(F.data.startswith("dl_rep:"))
async def handle_download_report_choice(callback: types.CallbackQuery, state: FSMContext):
    _, fmt, scope = callback.data.split(":")
    now = datetime.now()

    if scope == "weekly":
        sales = filter_sales_by_range(now - timedelta(days=7), now)
        title = "Haftalik_Z_Raporu"
    elif scope == "monthly":
        sales = filter_sales_by_range(datetime(now.year, now.month, 1), now)
        title = "Aylik_Z_Raporu"
    else:
        st_data = await state.get_data()
        sales = st_data.get("last_filtered_sales", [])
        title = st_data.get("report_title", "Z_Raporu")

    if not sales:
        await callback.answer("Bu dönem için onaylı satış kaydı bulunmuyor.", show_alert=True)
        return

    await callback.answer("Rapor dökümanı hazırlanıyor...")
    if fmt in ["excel", "both"]:
        f = generate_excel_z_report(sales, title)
        if f and os.path.exists(f):
            await callback.message.answer_document(FSInputFile(f), caption=f"📊 <b>Elips Depo {title} (Excel)</b>", parse_mode="HTML")

    if fmt in ["pdf", "both"]:
        f = generate_pdf_z_report(sales, title)
        if f and os.path.exists(f):
            await callback.message.answer_document(FSInputFile(f), caption=f"📑 <b>Elips Depo {title} (PDF)</b>", parse_mode="HTML")

# ----------------- TESİS GEÇMİŞİ -----------------
@dp.message(Command("gecmis", ignore_mention=True))
async def cmd_history_query(message: types.Message):
    parts = message.text.split(maxsplit=1)
    if len(parts) < 2:
        await message.answer("⚠️ Tesis adı belirtin. (Örn: <code>/gecmis TERRY JAR</code>)", parse_mode="HTML")
        return
    await message.answer(get_client_history_text(parts[1]), parse_mode="HTML")

@dp.message(F.text == "🏢 Tesis Geçmişi")
async def btn_client_history(message: types.Message, state: FSMContext):
    await state.set_state(ClientHistoryState.waiting_client_name)
    await message.answer("🔍 Geçmişini aradığınız <b>Fabrika Adını</b> yazın (Örn: Terry, Zelal, Asaka, Hajj):", parse_mode="HTML")

@dp.message(ClientHistoryState.waiting_client_name, ~F.text.startswith("/"))
async def process_client_history_input(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer(get_client_history_text(message.text.strip()), parse_mode="HTML")

# ----------------- WEBAPP VERİLERİ -----------------
@dp.message(F.web_app_data)
async def handle_webapp_payload(message: types.Message):
    try: p = json.loads(message.web_app_data.data)
    except Exception: return
    op = "Ömer" if message.chat.id == OMER_CHAT_ID else "Kamil"
    ref_id = f"REF-{message.message_id}"
    dir_type = p.get("direction")

    if dir_type == "STOCK_CHANGE" and p.get("action") == "DECREASE":
        dest = p.get("destination", "Bilinmeyen")
        code, name, qty = p.get("code"), p.get("name"), p.get("delta", 1)
        sales = load_sales()
        sales.append({
            "id": ref_id, "firm": dest, "source": "ELIPS_DEPO", "ref_no": "MANUEL-ÇIKIŞ",
            "items": [{"code": code, "name": name, "qty": qty}], "status": "APPROVED",
            "date": datetime.now().strftime("%d.%m.%Y %H:%M"), "timestamp": datetime.now().isoformat()
        })
        save_sales(sales)
        msg = f"📤 <b>PARÇA SATILDI / DEPO ÇIKIŞI!</b>\n🏢 Firma: <code>{html.escape(dest)}</code>\n⚙️ <code>[{code}]</code> {html.escape(name)}\n📉 Adet: <b>-{qty}</b> | Kalan: <code>{p.get('current_qty', 0)} ADT</code>\n👷‍♂️ Çıkış Yapan: {op}"
        await notify_both(msg)

# ----------------- SOHBET KÖPRÜSÜ (TELSİZ) -----------------
@dp.message(~F.text.startswith("/"))
async def chat_bridge(message: types.Message):
    if TEST_MODE:
        return

    menu_buttons = ["📦 Parça Listesi", "📊 Z-Raporu Al", "📅 Tarihli Rapor Al", "⏳ Bekleyen Teklifler", "✅ Tamamlanan Teklifler", "📂 Verilen Teklifler", "🏢 Tesis Geçmişi", "📝 Eliar'dan Teklif Al", "📖 Kataloğu Gönder"]
    if message.text in menu_buttons: return
    tid = KAMIL_CHAT_ID if message.chat.id == OMER_CHAT_ID else (OMER_CHAT_ID if message.chat.id == KAMIL_CHAT_ID else None)
    sender = "Ömer" if message.chat.id == OMER_CHAT_ID else "Kamil"
    if tid and message.text:
        await bot.send_message(tid, f"💬 <b>{sender}:</b> {html.escape(message.text)}", parse_mode="HTML")
        await message.reply("✓ <i>İletildi.</i>", parse_mode="HTML")

# ----------------- ZAMANLAYICI & SERVER -----------------
async def automated_scheduler():
    last_week = ""
    while True:
        now = datetime.now()
        today = now.strftime("%Y-%m-%d")
        if now.weekday() == 5 and now.hour == 12 and now.minute == 30 and last_week != today:
            last_week = today
            sales = filter_sales_by_range(now - timedelta(days=7), now)
            await notify_both("📢 <b>HAFTALIK OTOMATİK Z-RAPORU</b>\n\n" + build_report_text(sales, "HAFTALIK"))
            if sales:
                f_xls = generate_excel_z_report(sales, "Haftalik")
                f_pdf = generate_pdf_z_report(sales, "Haftalik")
                if f_xls: await bot.send_document(OMER_CHAT_ID, FSInputFile(f_xls))
                if f_pdf: await bot.send_document(OMER_CHAT_ID, FSInputFile(f_pdf))
        await asyncio.sleep(25)

async def start_dummy_server():
    try:
        app = web.Application()
        app.router.add_get('/', lambda r: web.Response(text="Elips Depo Robotu Aktif"))
        runner = web.AppRunner(app)
        await runner.setup()
        port = int(os.environ.get("PORT", 8080))
        await web.TCPSite(runner, '0.0.0.0', port).start()
    except Exception as e:
        logging.warning(f"Server başlatılamadı: {e}")

# ==============================================================================
#                               ANA ÇALIŞTIRICI
# ==============================================================================
async def main():
    try:
        # 170+ Kalem parçanın tamamını veritabanına zorunlu eşitle
        sync_all_stocks_to_db()

        await bot.delete_webhook(drop_pending_updates=True)
        await setup_bot_commands()
        await start_dummy_server()
        asyncio.create_task(automated_scheduler())
        
        mode_str = "🧪 TEST / SİMÜLASYON MODUNDA (Kamil kapalı, stoklar korunuyor)" if TEST_MODE else "🟢 CANLI MODDA (Kamil aktif)"
        logging.info(f">>> ELIPS DEPO BOTU BAŞLATILDI: {mode_str} <<<")
        await dp.start_polling(bot)
    except Exception as e:
        logging.error(f"Kritik hata: {e}")

if __name__ == "__main__":
    asyncio.run(main())
